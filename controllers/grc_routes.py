"""
Governance, Risk & Compliance (GRC) Module - Routes and Controllers
===================================================================
Enterprise-grade GRC module for the MMDx ERP system covering:
- Governance Structure Management
- Risk Register & Assessment
- Control Library & Testing
- Compliance Obligations & Regulations
- Policy Management & Acknowledgements
- Segregation of Duties (SoD)
- Access Reviews
- Violations, Exceptions & Findings
- Remediation & CAPA Plans
- Evidence Repository
- Audit Readiness Center
- Incidents & Breach Tracking
- Certifications & Renewals
- Alert Rules & Events
- High-Risk Activity Monitoring
- Approval Matrix Governance
- Reports & Analytics

Author: Enterprise Architecture Team
Version: 1.0.0
"""

import functools
import json
import math
import re
import uuid
import csv
import io
from datetime import datetime, timedelta
from typing import Optional, List, Dict, Any

from flask import Flask, request, jsonify, render_template, session, send_file, redirect, url_for, flash, Response
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

import database as db_helper


def get_grc_db():
    """Get database connection for GRC module."""
    return db_helper.get_db()


def init_grc_module(app: Flask):
    """
    Initialize the Governance & Compliance (GRC) module.
    Called during app startup via app.py register mechanism.
    """
    from grc_models import init_grc_tables, seed_grc_initial_data
    import database as db_helper

    conn = db_helper.get_db()
    cursor = conn.cursor()

    init_grc_tables()
    seed_grc_initial_data()

    print("[GRC Module] Governance & Compliance initialized successfully")


def require_grc_permission(resource: str, action: str = 'view'):
    """Permission decorator for GRC routes."""
    def decorator(f):
        @functools.wraps(f)
        def decorated_function(*args, **kwargs):
            from permissions import user_has_permission
            user_id = session.get('user_id')
            if not user_id:
                if request.is_json:
                    return jsonify({'error': 'Authentication required'}), 401
                return redirect(url_for('login'))

            if not user_has_permission(user_id, 'grc', resource, action):
                if request.is_json:
                    return jsonify({'error': 'Permission denied'}), 403
                flash('You do not have permission to perform this action.', 'error')
                return redirect(url_for('dashboard'))

            return f(*args, **kwargs)
        return decorated_function
    return decorator


def require_login(f):
    """Require user to be logged in."""
    @functools.wraps(f)
    def decorated_function(*args, **kwargs):
        if 'user_id' not in session:
            if request.is_json:
                return jsonify({'error': 'Authentication required'}), 401
            return redirect(url_for('login'))
        return f(*args, **kwargs)
    return decorated_function


def get_current_user_id():
    """Get current user ID from session."""
    return session.get('user_id')


def get_current_company_id():
    """Get current company ID from session."""
    return session.get('company_id')


def parse_sort_params(default_sort='created_at', default_order='desc'):
    """Parse sort parameters from request."""
    sort = request.args.get('sort', default_sort)
    order = request.args.get('order', default_order)
    if order not in ('asc', 'desc'):
        order = default_order
    return sort, order


def parse_pagination_params(default_page=1, default_per_page=20):
    """Parse pagination parameters from request."""
    page = max(1, int(request.args.get('page', default_page)))
    per_page = min(100, max(10, int(request.args.get('per_page', default_per_page))))
    return page, per_page


def build_pagination_response(items: List, total: int, page: int, per_page: int, **kwargs):
    """Build pagination metadata for responses."""
    total_pages = math.ceil(total / per_page) if per_page > 0 else 0
    return {
        'items': items,
        'pagination': {
            'page': page,
            'per_page': per_page,
            'total': total,
            'total_pages': total_pages,
            'has_next': page < total_pages,
            'has_prev': page > 1
        },
        **kwargs
    }


def log_grc_audit(entity_type: str, entity_id: int, action: str,
                  user_id: int = None, notes: str = None,
                  old_value: str = None, new_value: str = None):
    """Log audit entry for GRC module."""
    from database import log_audit
    company_id = get_current_company_id()
    log_audit(
        entity_type=entity_type,
        entity_id=entity_id,
        action=action,
        user_id=user_id or get_current_user_id(),
        notes=notes,
        old_value=old_value,
        new_value=new_value,
        company_id=company_id
    )


def notify_flow_channel(channel_name: str, message: str, severity: str = 'info'):
    """Send notification to Flow channel if available."""
    try:
        from flow_models import send_channel_message
        company_id = get_current_company_id()
        send_channel_message(channel_name, message, company_id)
    except Exception as e:
        print(f"[Flow Notification] Failed to send to {channel_name}: {e}")


def generate_grc_code(prefix: str) -> str:
    """Generate unique GRC code."""
    year = datetime.now().strftime('%Y')
    return f"GRC-{prefix}-{year}-{uuid.uuid4().hex[:6].upper()}"


def export_to_excel_text(data: List[Dict], columns: List[str], title: str) -> Response:
    """Export data to Excel with all cells as text."""
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]

    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    for col_idx, col_name in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for row_idx, row in enumerate(data, 2):
        for col_idx, col_name in enumerate(columns, 1):
            value = row.get(col_name, '')
            cell = ws.cell(row=row_idx, column=col_idx, value=str(value) if value is not None else '')
            cell.border = border
            cell.alignment = Alignment(horizontal='left', vertical='center')

    for col_idx, col_name in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = max(15, len(col_name) + 2)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=f"{title}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    )


def export_to_excel_general(data: List[Dict], columns: List[str], title: str) -> Response:
    """Export data to Excel with cells as General format."""
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]

    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    for col_idx, col_name in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for row_idx, row in enumerate(data, 2):
        for col_idx, col_name in enumerate(columns, 1):
            value = row.get(col_name, '')
            if value is None:
                value = ''
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = border
            cell.alignment = Alignment(horizontal='left', vertical='center')

    for col_idx, col_name in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = max(15, len(col_name) + 2)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=f"{title}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    )


def export_to_csv(data: List[Dict], columns: List[str], title: str) -> Response:
    """Export data to CSV format."""
    output = io.StringIO()
    writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL)

    # Write header
    writer.writerow(columns)

    # Write data rows
    for row in data:
        row_values = []
        for col in columns:
            value = row.get(col, '')
            if value is None:
                value = ''
            row_values.append(str(value))
        writer.writerow(row_values)

    output.seek(0)
    return send_file(
        output,
        mimetype='text/csv',
        as_attachment=True,
        download_name=f"{title}_{datetime.now().strftime('%Y%m%d')}.csv"
    )


# =====================================================================
# GRC ROUTES REGISTRATION
# =====================================================================

def register_grc_routes(app: Flask):
    """
    Register all GRC routes. This is the main entry point called from app.py.
    """

    # -------------------------------------------------------------------------
    # DASHBOARD ROUTES
    # -------------------------------------------------------------------------

    @app.route('/grc')
    @app.route('/grc/dashboard')
    @require_login
    def grc_dashboard():
        """Main GRC dashboard."""
        company_id = get_current_company_id()
        user_id = get_current_user_id()
        conn = get_grc_db()
        cursor = conn.cursor()

        stats = {}
        recent_items = {}

        cursor.execute("SELECT COUNT(*) as cnt FROM grc_risks WHERE status = 'active'")
        stats['total_risks'] = cursor.fetchone()['cnt']

        cursor.execute("SELECT COUNT(*) as cnt FROM grc_risks WHERE status = 'active' AND risk_level = 'high'")
        stats['high_risks'] = cursor.fetchone()['cnt']

        cursor.execute("SELECT COUNT(*) as cnt FROM grc_controls WHERE status = 'active'")
        stats['total_controls'] = cursor.fetchone()['cnt']

        cursor.execute("""
            SELECT COUNT(*) as cnt FROM grc_control_tests
            WHERE status IN ('in_progress', 'pending') AND next_test_date < date('now')
        """)
        stats['overdue_tests'] = cursor.fetchone()['cnt']

        cursor.execute("SELECT COUNT(*) as cnt FROM grc_findings WHERE status = 'open'")
        stats['open_findings'] = cursor.fetchone()['cnt']

        cursor.execute("SELECT COUNT(*) as cnt FROM grc_sod_conflicts WHERE status = 'open'")
        stats['sod_conflicts'] = cursor.fetchone()['cnt']

        cursor.execute("SELECT COUNT(*) as cnt FROM grc_exceptions WHERE status IN ('open', 'approved')")
        stats['open_exceptions'] = cursor.fetchone()['cnt']

        cursor.execute("""
            SELECT COUNT(*) as cnt FROM grc_certifications
            WHERE renewal_status = 'valid' AND expiry_date BETWEEN date('now') AND date('now', '+30 days')
        """)
        stats['expiring_certs'] = cursor.fetchone()['cnt']

        cursor.execute("SELECT COUNT(*) as cnt FROM grc_policies WHERE acknowledgement_required = 1")
        stats['total_policies'] = cursor.fetchone()['cnt']

        cursor.execute("""
            SELECT COUNT(*) as cnt FROM grc_incidents
            WHERE status = 'open'
        """)
        stats['open_incidents'] = cursor.fetchone()['cnt']

        cursor.execute("""
            SELECT * FROM grc_risks
            WHERE status = 'active'
            ORDER BY inherent_risk_score DESC
            LIMIT 5
        """)
        stats['top_risks'] = [dict(row) for row in cursor.fetchall()]

        cursor.execute("""
            SELECT * FROM grc_findings
            WHERE status = 'open'
            ORDER BY created_at DESC
            LIMIT 5
        """)
        stats['recent_findings'] = [dict(row) for row in cursor.fetchall()]

        cursor.execute("""
            SELECT * FROM grc_control_tests
            WHERE status = 'pending' OR (status = 'in_progress' AND next_test_date < date('now'))
            ORDER BY next_test_date ASC
            LIMIT 5
        """)
        stats['pending_tests'] = [dict(row) for row in cursor.fetchall()]

        conn.close()

        return render_template('grc/dashboard/index.html',
                           stats=stats,
                           recent_items=recent_items)

    @app.route('/grc/dashboard/executive')
    @require_login
    @require_grc_permission('dashboard', 'view')
    def grc_executive_dashboard():
        """Executive GRC dashboard with KPIs."""
        company_id = get_current_company_id()
        conn = get_grc_db()
        cursor = conn.cursor()

        kpis = {}

        cursor.execute("SELECT COUNT(*) as cnt FROM grc_risks WHERE status = 'active'")
        kpis['total_risks'] = cursor.fetchone()['cnt']

        cursor.execute("""
            SELECT COUNT(*) as cnt FROM grc_risks
            WHERE status = 'active' AND inherent_risk_score >= 6
        """)
        kpis['high_risks'] = cursor.fetchone()['cnt']

        cursor.execute("""
            SELECT COUNT(*) as cnt FROM grc_controls
            WHERE status = 'active' AND operating_effectiveness = 1
        """)
        kpis['effective_controls'] = cursor.fetchone()['cnt']

        cursor.execute("""
            SELECT COUNT(*) as cnt FROM grc_control_tests
            WHERE status = 'completed' AND effectiveness_rating >= 3
        """)
        kpis['passed_tests'] = cursor.fetchone()['cnt']

        cursor.execute("SELECT COUNT(*) as cnt FROM grc_sod_conflicts WHERE status = 'open'")
        kpis['sod_conflicts'] = cursor.fetchone()['cnt']

        cursor.execute("SELECT COUNT(*) as cnt FROM grc_findings WHERE status = 'open'")
        kpis['open_findings'] = cursor.fetchone()['cnt']

        cursor.execute("SELECT COUNT(*) as cnt FROM grc_exceptions WHERE status = 'open'")
        kpis['open_exceptions'] = cursor.fetchone()['cnt']

        cursor.execute("""
            SELECT COUNT(*) as cnt FROM grc_policy_acknowledgements
            WHERE status = 'pending'
        """)
        kpis['pending_acks'] = cursor.fetchone()['cnt']

        conn.close()

        return render_template('grc/dashboard/executive.html', kpis=kpis)

    # -------------------------------------------------------------------------
    # GOVERNANCE STRUCTURE ROUTES
    # -------------------------------------------------------------------------

    @app.route('/grc/governance')
    @require_login
    @require_grc_permission('governance', 'view')
    def grc_governance_list():
        """List all governance entities."""
        company_id = get_current_company_id()
        conn = get_grc_db()
        cursor = conn.cursor()

        cursor.execute("""
            SELECT ge.*, u1.full_name as owner_name, u2.full_name as compliance_officer_name
            FROM grc_governance_entities ge
            LEFT JOIN users u1 ON ge.owner_user_id = u1.id
            LEFT JOIN users u2 ON ge.compliance_officer_id = u2.id
            WHERE ge.is_active = 1
            ORDER BY ge.entity_name
        """)
        entities = [dict(row) for row in cursor.fetchall()]
        conn.close()

        return render_template('grc/governance/index.html', entities=entities)

    @app.route('/grc/governance/create', methods=['GET', 'POST'])
    @require_login
    @require_grc_permission('governance', 'create')
    def grc_governance_create():
        """Create new governance entity."""
        if request.method == 'POST':
            conn = get_grc_db()
            cursor = conn.cursor()

            cursor.execute("""
                INSERT INTO grc_governance_entities (
                    entity_code, entity_name, entity_type, description,
                    company_id, branch_id, department_id, process_area,
                    owner_user_id, compliance_officer_id, risk_owner_id,
                    control_owner_id, status, created_at, created_by
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'active', datetime('now'), ?)
            """, (
                generate_grc_code('GOV'),
                request.form.get('entity_name'),
                request.form.get('entity_type'),
                request.form.get('description'),
                get_current_company_id(),
                request.form.get('branch_id'),
                request.form.get('department_id'),
                request.form.get('process_area'),
                request.form.get('owner_user_id'),
                request.form.get('compliance_officer_id'),
                request.form.get('risk_owner_id'),
                request.form.get('control_owner_id'),
                get_current_user_id()
            ))
            entity_id = cursor.lastrowid
            conn.commit()

            log_grc_audit('governance_entity', entity_id, 'CREATE',
                         notes=f"Created governance entity: {request.form.get('entity_name')}")

            conn.close()
            flash('Governance entity created successfully.', 'success')
            return redirect(url_for('grc_governance_list'))

        return render_template('grc/governance/create.html')

    @app.route('/grc/risks')
    @require_login
    @require_grc_permission('risks', 'view')
    def grc_risks_list():
        """List all risks."""
        company_id = get_current_company_id()
        conn = get_grc_db()
        cursor = conn.cursor()

        page, per_page = parse_pagination_params()
        sort, order = parse_sort_params()

        cursor.execute("SELECT COUNT(*) as cnt FROM grc_risks WHERE status = 'active'")
        total = cursor.fetchone()['cnt']

        offset = (page - 1) * per_page
        cursor.execute(f"""
            SELECT r.*, rc.category_name, u.full_name as owner_name
            FROM grc_risks r
            LEFT JOIN grc_risk_categories rc ON r.category_id = rc.id
            LEFT JOIN users u ON r.risk_owner_id = u.id
            WHERE r.status = 'active'
            ORDER BY r.{sort} {order}
            LIMIT ? OFFSET ?
        """, (per_page, offset))
        risks = [dict(row) for row in cursor.fetchall()]

        conn.close()

        return render_template('grc/risks/index.html',
                             risks=risks,
                             page=page,
                             per_page=per_page,
                             total=total)

    @app.route('/grc/risks/assessments')
    @require_login
    @require_grc_permission('risk_assessments', 'view')
    def grc_risk_assessments():
        """Risk assessments view - reuses risk register with assessment filters."""
        conn = get_grc_db()
        cursor = conn.cursor()

        page, per_page = parse_pagination_params()
        sort, order = parse_sort_params()

        # Filter for risks that need assessment review
        cursor.execute("""
            SELECT COUNT(*) as cnt FROM grc_risks
            WHERE status = 'active'
            AND (next_review_date IS NULL OR next_review_date <= date('now', '+30 days'))
        """)
        total = cursor.fetchone()['cnt']

        offset = (page - 1) * per_page
        cursor.execute(f"""
            SELECT r.*,
                   rc.category_name,
                   u.full_name as owner_name,
                   u2.full_name as reviewer_name
            FROM grc_risks r
            LEFT JOIN grc_risk_categories rc ON r.category_id = rc.id
            LEFT JOIN users u ON r.risk_owner_id = u.id
            LEFT JOIN users u2 ON r.reviewer_user_id = u2.id
            WHERE r.status = 'active'
            AND (r.next_review_date IS NULL OR r.next_review_date <= date('now', '+30 days'))
            ORDER BY r.{sort} {order}
            LIMIT ? OFFSET ?
        """, (per_page, offset))
        risks = [dict(row) for row in cursor.fetchall()]

        conn.close()

        return render_template('grc/risks/index.html',
                             risks=risks,
                             page=page,
                             per_page=per_page,
                             total=total,
                             view_title='Risk Assessments')

    @app.route('/grc/risks/create', methods=['GET', 'POST'])
    @require_login
    @require_grc_permission('risks', 'create')
    def grc_risk_create():
        """Create new risk."""
        conn = get_grc_db()
        cursor = conn.cursor()

        cursor.execute("SELECT * FROM grc_risk_categories WHERE is_active = 1")
        categories = [dict(row) for row in cursor.fetchall()]

        if request.method == 'POST':
            impact = int(request.form.get('impact_score', 3))
            likelihood = int(request.form.get('likelihood_score', 3))
            inherent_risk = impact * likelihood

            cursor.execute("""
                INSERT INTO grc_risks (
                    risk_code, title, description, category_id, risk_type,
                    company_id, department_id, business_unit, process_area,
                    risk_owner_id, impact_score, likelihood_score,
                    inherent_risk_score, risk_level, status,
                    trigger_conditions, leading_indicators, due_date,
                    treatment_plan, created_at, created_by
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'active', ?, ?, ?, ?, ?, datetime('now'), ?)
            """, (
                generate_grc_code('RSK'),
                request.form.get('title'),
                request.form.get('description'),
                request.form.get('category_id'),
                request.form.get('risk_type'),
                get_current_company_id(),
                request.form.get('department_id'),
                request.form.get('business_unit'),
                request.form.get('process_area'),
                request.form.get('risk_owner_id'),
                impact,
                likelihood,
                inherent_risk,
                'high' if inherent_risk >= 6 else 'medium' if inherent_risk >= 4 else 'low',
                request.form.get('trigger_conditions'),
                request.form.get('leading_indicators'),
                request.form.get('due_date'),
                request.form.get('treatment_plan'),
                get_current_user_id()
            ))
            risk_id = cursor.lastrowid
            conn.commit()

            log_grc_audit('risk', risk_id, 'CREATE',
                         notes=f"Created risk: {request.form.get('title')}")

            conn.close()
            flash('Risk created successfully.', 'success')
            return redirect(url_for('grc_risks_list'))

        conn.close()
        return render_template('grc/risks/create.html', categories=categories)

    @app.route('/grc/controls')
    @require_login
    @require_grc_permission('controls', 'view')
    def grc_controls_list():
        """List all controls."""
        company_id = get_current_company_id()
        conn = get_grc_db()
        cursor = conn.cursor()

        cursor.execute("""
            SELECT c.*, cc.category_name, u.full_name as owner_name
            FROM grc_controls c
            LEFT JOIN grc_control_categories cc ON c.category_id = cc.id
            LEFT JOIN users u ON c.owner_user_id = u.id
            WHERE c.is_active = 1
            ORDER BY c.control_code
        """)
        controls = [dict(row) for row in cursor.fetchall()]

        conn.close()
        return render_template('grc/controls/index.html', controls=controls)

    @app.route('/grc/controls/create', methods=['GET', 'POST'])
    @require_login
    @require_grc_permission('controls', 'create')
    def grc_control_create():
        """Create new control."""
        if request.method == 'POST':
            conn = get_grc_db()
            cursor = conn.cursor()

            cursor.execute("""
                INSERT INTO grc_controls (
                    control_code, title, objective, description,
                    control_type, nature, frequency, category_id,
                    company_id, department_id, owner_user_id,
                    evidence_requirement, test_method, status,
                    key_control, financial_control, access_control,
                    inventory_control, customs_compliance,
                    created_at, created_by
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'active', ?, ?, ?, ?, ?, ?, datetime('now'), ?)
            """, (
                generate_grc_code('CTL'),
                request.form.get('title'),
                request.form.get('objective'),
                request.form.get('description'),
                request.form.get('control_type'),
                request.form.get('nature'),
                request.form.get('frequency'),
                request.form.get('category_id'),
                get_current_company_id(),
                request.form.get('department_id'),
                request.form.get('owner_user_id'),
                request.form.get('evidence_requirement'),
                request.form.get('test_method'),
                1 if request.form.get('key_control') else 0,
                1 if request.form.get('financial_control') else 0,
                1 if request.form.get('access_control') else 0,
                1 if request.form.get('inventory_control') else 0,
                1 if request.form.get('customs_compliance') else 0,
                get_current_user_id()
            ))
            control_id = cursor.lastrowid
            conn.commit()

            log_grc_audit('control', control_id, 'CREATE',
                         notes=f"Created control: {request.form.get('title')}")

            conn.close()
            flash('Control created successfully.', 'success')
            return redirect(url_for('grc_controls_list'))

        conn = get_grc_db()
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM grc_control_categories WHERE is_active = 1")
        categories = [dict(row) for row in cursor.fetchall()]
        conn.close()

        return render_template('grc/controls/create.html', categories=categories)

    @app.route('/grc/controls/tests')
    @require_login
    @require_grc_permission('control_tests', 'view')
    def grc_control_tests_list():
        """List all control tests."""
        conn = get_grc_db()
        cursor = conn.cursor()

        cursor.execute("""
            SELECT ct.*, c.title as control_title, c.control_code,
                   u1.full_name as tester_name, u2.full_name as reviewer_name
            FROM grc_control_tests ct
            JOIN grc_controls c ON ct.control_id = c.id
            LEFT JOIN users u1 ON ct.tester_user_id = u1.id
            LEFT JOIN users u2 ON ct.reviewer_user_id = u2.id
            ORDER BY ct.test_date DESC
        """)
        tests = [dict(row) for row in cursor.fetchall()]

        conn.close()
        return render_template('grc/controls/tests.html', tests=tests)

    @app.route('/grc/controls/tests/create', methods=['GET', 'POST'])
    @require_login
    @require_grc_permission('control_tests', 'create')
    def grc_control_test_create():
        """Create new control test."""
        conn = get_grc_db()
        cursor = conn.cursor()

        cursor.execute("SELECT * FROM grc_controls WHERE is_active = 1")
        controls = [dict(row) for row in cursor.fetchall()]

        if request.method == 'POST':
            cursor.execute("""
                INSERT INTO grc_control_tests (
                    test_code, control_id, test_date, tester_user_id,
                    sample_size, test_procedure, status,
                    next_test_date, created_at, created_by
                ) VALUES (?, ?, ?, ?, ?, ?, 'in_progress', ?, datetime('now'), ?)
            """, (
                generate_grc_code('TST'),
                request.form.get('control_id'),
                request.form.get('test_date'),
                get_current_user_id(),
                request.form.get('sample_size'),
                request.form.get('test_procedure'),
                request.form.get('next_test_date'),
                get_current_user_id()
            ))
            test_id = cursor.lastrowid
            conn.commit()

            log_grc_audit('control_test', test_id, 'CREATE',
                         notes=f"Created control test for control ID: {request.form.get('control_id')}")

            conn.close()
            flash('Control test created successfully.', 'success')
            return redirect(url_for('grc_control_tests_list'))

        conn.close()
        return render_template('grc/controls/test_create.html', controls=controls)

    @app.route('/grc/obligations')
    @require_login
    @require_grc_permission('obligations', 'view')
    def grc_obligations_list():
        """List all compliance obligations."""
        conn = get_grc_db()
        cursor = conn.cursor()

        cursor.execute("""
            SELECT o.*, u.full_name as owner_name
            FROM grc_compliance_obligations o
            LEFT JOIN users u ON o.owner_user_id = u.id
            ORDER BY o.obligation_code
        """)
        obligations = [dict(row) for row in cursor.fetchall()]

        conn.close()
        return render_template('grc/obligations/index.html', obligations=obligations)

    @app.route('/grc/regulations')
    @require_login
    @require_grc_permission('regulations', 'view')
    def grc_regulations_list():
        """List all regulations."""
        conn = get_grc_db()
        cursor = conn.cursor()

        cursor.execute("SELECT * FROM grc_regulations ORDER BY regulation_name")
        regulations = [dict(row) for row in cursor.fetchall()]

        conn.close()
        return render_template('grc/regulations/index.html', regulations=regulations)

    @app.route('/grc/policies')
    @require_login
    @require_grc_permission('policies', 'view')
    def grc_policies_list():
        """List all policies."""
        conn = get_grc_db()
        cursor = conn.cursor()

        cursor.execute("""
            SELECT p.*, u.full_name as owner_name
            FROM grc_policy_documents p
            LEFT JOIN users u ON p.owner_user_id = u.id
            ORDER BY p.title
        """)
        policies = [dict(row) for row in cursor.fetchall()]

        conn.close()
        return render_template('grc/policies/index.html', policies=policies)

    @app.route('/grc/policies/acknowledgements')
    @require_login
    @require_grc_permission('policy_acknowledgements', 'view')
    def grc_policy_acknowledgements_list():
        """List policy acknowledgements."""
        conn = get_grc_db()
        cursor = conn.cursor()

        cursor.execute("""
            SELECT pa.*, p.title as policy_title, u.full_name as user_name
            FROM grc_policy_acknowledgements pa
            JOIN grc_policy_documents p ON pa.policy_id = p.id
            LEFT JOIN users u ON pa.user_id = u.id
            ORDER BY pa.due_date
        """)
        acknowledgements = [dict(row) for row in cursor.fetchall()]

        conn.close()
        return render_template('grc/policies/acknowledgements.html', acknowledgements=acknowledgements)

    @app.route('/grc/sod')
    @require_login
    @require_grc_permission('sod', 'view')
    def grc_sod_list():
        """List all SoD rules."""
        conn = get_grc_db()
        cursor = conn.cursor()

        cursor.execute("SELECT * FROM grc_sod_rules WHERE is_active = 1 ORDER BY rule_name")
        rules = [dict(row) for row in cursor.fetchall()]

        conn.close()
        return render_template('grc/sod/index.html', rules=rules)

    @app.route('/grc/sod/conflicts')
    @require_login
    @require_grc_permission('sod_conflicts', 'view')
    def grc_sod_conflicts_list():
        """List all SoD conflicts."""
        conn = get_grc_db()
        cursor = conn.cursor()

        cursor.execute("""
            SELECT sc.*, sr.rule_name, u.full_name as user_name
            FROM grc_sod_conflicts sc
            JOIN grc_sod_rules sr ON sc.rule_id = sr.id
            LEFT JOIN users u ON sc.user_id = u.id
            ORDER BY sc.detected_at DESC
        """)
        conflicts = [dict(row) for row in cursor.fetchall()]

        conn.close()
        return render_template('grc/sod/conflicts.html', conflicts=conflicts)

    @app.route('/grc/access-reviews')
    @require_login
    @require_grc_permission('access_reviews', 'view')
    def grc_access_reviews_list():
        """List all access reviews."""
        conn = get_grc_db()
        cursor = conn.cursor()

        cursor.execute("""
            SELECT ar.*, u.full_name as reviewer_name
            FROM grc_access_reviews ar
            LEFT JOIN users u ON ar.reviewer_user_id = u.id
            ORDER BY ar.created_at DESC
        """)
        reviews = [dict(row) for row in cursor.fetchall()]

        conn.close()
        return render_template('grc/access_reviews/index.html', reviews=reviews)

    @app.route('/grc/exceptions')
    @require_login
    @require_grc_permission('exceptions', 'view')
    def grc_exceptions_list():
        """List all exceptions."""
        conn = get_grc_db()
        cursor = conn.cursor()

        cursor.execute("""
            SELECT e.*, u.full_name as owner_name
            FROM grc_exceptions e
            LEFT JOIN users u ON e.owner_user_id = u.id
            ORDER BY e.created_at DESC
        """)
        exceptions = [dict(row) for row in cursor.fetchall()]

        conn.close()
        return render_template('grc/exceptions/index.html', exceptions=exceptions)

    @app.route('/grc/exceptions/create', methods=['GET', 'POST'])
    @require_login
    @require_grc_permission('exceptions', 'create')
    def grc_exception_create():
        """Create new exception."""
        if request.method == 'POST':
            conn = get_grc_db()
            cursor = conn.cursor()

            cursor.execute("""
                INSERT INTO grc_exceptions (
                    exception_code, title, description, exception_type,
                    severity, source, owner_user_id, status,
                    business_justification, compensating_controls,
                    start_date, expiry_date, created_at, created_by
                ) VALUES (?, ?, ?, ?, ?, ?, ?, 'open', ?, ?, ?, ?, datetime('now'), ?)
            """, (
                generate_grc_code('EXP'),
                request.form.get('title'),
                request.form.get('description'),
                request.form.get('exception_type'),
                request.form.get('severity'),
                request.form.get('source'),
                get_current_user_id(),
                request.form.get('business_justification'),
                request.form.get('compensating_controls'),
                request.form.get('start_date'),
                request.form.get('expiry_date'),
                get_current_user_id()
            ))
            exception_id = cursor.lastrowid
            conn.commit()

            log_grc_audit('exception', exception_id, 'CREATE',
                         notes=f"Created exception: {request.form.get('title')}")

            conn.close()
            flash('Exception created successfully.', 'success')
            return redirect(url_for('grc_exceptions_list'))

        return render_template('grc/exceptions/create.html')

    @app.route('/grc/findings')
    @require_login
    @require_grc_permission('findings', 'view')
    def grc_findings_list():
        """List all findings."""
        conn = get_grc_db()
        cursor = conn.cursor()

        cursor.execute("""
            SELECT f.*, u.full_name as owner_name
            FROM grc_findings f
            LEFT JOIN users u ON f.owner_user_id = u.id
            ORDER BY f.created_at DESC
        """)
        findings = [dict(row) for row in cursor.fetchall()]

        conn.close()
        return render_template('grc/findings/index.html', findings=findings)

    @app.route('/grc/findings/create', methods=['GET', 'POST'])
    @require_login
    @require_grc_permission('findings', 'create')
    def grc_finding_create():
        """Create new finding."""
        if request.method == 'POST':
            conn = get_grc_db()
            cursor = conn.cursor()

            cursor.execute("""
                INSERT INTO grc_findings (
                    finding_code, title, description, source, severity,
                    root_cause, impact, recommendation,
                    owner_user_id, status, target_closure_date,
                    created_at, created_by
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 'open', ?, datetime('now'), ?)
            """, (
                generate_grc_code('FND'),
                request.form.get('title'),
                request.form.get('description'),
                request.form.get('source'),
                request.form.get('severity'),
                request.form.get('root_cause'),
                request.form.get('impact'),
                request.form.get('recommendation'),
                request.form.get('owner_user_id'),
                request.form.get('target_closure_date'),
                get_current_user_id()
            ))
            finding_id = cursor.lastrowid
            conn.commit()

            log_grc_audit('finding', finding_id, 'CREATE',
                         notes=f"Created finding: {request.form.get('title')}")

            conn.close()
            flash('Finding created successfully.', 'success')
            return redirect(url_for('grc_findings_list'))

        return render_template('grc/findings/create.html')

    @app.route('/grc/remediation')
    @require_login
    @require_grc_permission('remediation', 'view')
    def grc_remediation_list():
        """List all remediation plans."""
        conn = get_grc_db()
        cursor = conn.cursor()

        cursor.execute("""
            SELECT rp.*, f.title as finding_title, u.full_name as owner_name
            FROM grc_remediation_plans rp
            LEFT JOIN grc_findings f ON rp.finding_id = f.id
            LEFT JOIN users u ON rp.owner_user_id = u.id
            ORDER BY rp.created_at DESC
        """)
        plans = [dict(row) for row in cursor.fetchall()]

        conn.close()
        return render_template('grc/remediation/index.html', plans=plans)

    @app.route('/grc/remediation/create', methods=['GET', 'POST'])
    @require_login
    @require_grc_permission('remediation', 'create')
    def grc_remediation_create():
        """Create new remediation plan."""
        if request.method == 'POST':
            conn = get_grc_db()
            cursor = conn.cursor()

            cursor.execute("""
                INSERT INTO grc_remediation_plans (
                    plan_code, title, description, finding_id,
                    owner_user_id, status, priority, progress_percent,
                    start_date, target_date, created_at, created_by
                ) VALUES (?, ?, ?, ?, ?, 'draft', ?, 0, ?, ?, datetime('now'), ?)
            """, (
                generate_grc_code('REM'),
                request.form.get('title'),
                request.form.get('description'),
                request.form.get('finding_id'),
                get_current_user_id(),
                request.form.get('priority'),
                request.form.get('start_date'),
                request.form.get('target_date'),
                get_current_user_id()
            ))
            plan_id = cursor.lastrowid
            conn.commit()

            log_grc_audit('remediation_plan', plan_id, 'CREATE',
                         notes=f"Created remediation plan: {request.form.get('title')}")

            conn.close()
            flash('Remediation plan created successfully.', 'success')
            return redirect(url_for('grc_remediation_list'))

        conn = get_grc_db()
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM grc_findings WHERE status = 'open'")
        findings = [dict(row) for row in cursor.fetchall()]
        conn.close()

        return render_template('grc/remediation/create.html', findings=findings)

    @app.route('/grc/evidence')
    @require_login
    @require_grc_permission('evidence', 'view')
    def grc_evidence_list():
        """List all evidence items."""
        conn = get_grc_db()
        cursor = conn.cursor()

        cursor.execute("""
            SELECT e.*, u.full_name as uploader_name
            FROM grc_evidence_items e
            LEFT JOIN users u ON e.uploader_user_id = u.id
            ORDER BY e.created_at DESC
        """)
        evidence_items = [dict(row) for row in cursor.fetchall()]

        conn.close()
        return render_template('grc/evidence/index.html', evidence_items=evidence_items)

    @app.route('/grc/audit')
    @require_login
    @require_grc_permission('audit_preparedness', 'view')
    def grc_audit_readiness():
        """Audit readiness center."""
        conn = get_grc_db()
        cursor = conn.cursor()

        cursor.execute("""
            SELECT * FROM grc_audit_preparedness_checks
            ORDER BY created_at DESC
        """)
        checks = [dict(row) for row in cursor.fetchall()]

        conn.close()
        return render_template('grc/audit/index.html', checks=checks)

    @app.route('/grc/incidents')
    @require_login
    @require_grc_permission('incidents', 'view')
    def grc_incidents_list():
        """List all incidents."""
        conn = get_grc_db()
        cursor = conn.cursor()

        cursor.execute("""
            SELECT i.*, u.full_name as discovered_by_name
            FROM grc_incidents i
            LEFT JOIN users u ON i.discovered_by = u.id
            ORDER BY i.created_at DESC
        """)
        incidents = [dict(row) for row in cursor.fetchall()]

        conn.close()
        return render_template('grc/incidents/index.html', incidents=incidents)

    @app.route('/grc/incidents/create', methods=['GET', 'POST'])
    @require_login
    @require_grc_permission('incidents', 'create')
    def grc_incident_create():
        """Create new incident."""
        if request.method == 'POST':
            conn = get_grc_db()
            cursor = conn.cursor()

            cursor.execute("""
                INSERT INTO grc_incidents (
                    incident_code, title, description, category, source,
                    incident_datetime, discovered_by, affected_module,
                    severity, status, created_at, created_by
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 'open', datetime('now'), ?)
            """, (
                generate_grc_code('INC'),
                request.form.get('title'),
                request.form.get('description'),
                request.form.get('category'),
                request.form.get('source'),
                request.form.get('incident_datetime'),
                get_current_user_id(),
                request.form.get('affected_module'),
                request.form.get('severity'),
                get_current_user_id()
            ))
            incident_id = cursor.lastrowid
            conn.commit()

            log_grc_audit('incident', incident_id, 'CREATE',
                         notes=f"Created incident: {request.form.get('title')}")

            notify_flow_channel('grc-alerts',
                              f"New GRC Incident: {request.form.get('title')} - Severity: {request.form.get('severity')}",
                              'warning')

            conn.close()
            flash('Incident created successfully.', 'success')
            return redirect(url_for('grc_incidents_list'))

        return render_template('grc/incidents/create.html')

    @app.route('/grc/certifications')
    @require_login
    @require_grc_permission('certifications', 'view')
    def grc_certifications_list():
        """List all certifications."""
        conn = get_grc_db()
        cursor = conn.cursor()

        cursor.execute("""
            SELECT c.*, u.full_name as owner_name
            FROM grc_certifications c
            LEFT JOIN users u ON c.owner_user_id = u.id
            ORDER BY c.expiry_date
        """)
        certifications = [dict(row) for row in cursor.fetchall()]

        conn.close()
        return render_template('grc/certifications/index.html', certifications=certifications)

    @app.route('/grc/deadlines')
    @require_login
    @require_grc_permission('deadlines', 'view')
    def grc_deadlines_list():
        """List all deadlines."""
        conn = get_grc_db()
        cursor = conn.cursor()

        cursor.execute("""
            SELECT d.*, u.full_name as owner_name
            FROM grc_deadlines d
            LEFT JOIN users u ON d.owner_user_id = u.id
            WHERE d.status = 'open'
            ORDER BY d.due_date
        """)
        deadlines = [dict(row) for row in cursor.fetchall()]

        conn.close()
        return render_template('grc/deadlines/index.html', deadlines=deadlines)

    @app.route('/grc/alerts/rules')
    @require_login
    @require_grc_permission('alert_rules', 'view')
    def grc_alert_rules_list():
        """List all alert rules."""
        conn = get_grc_db()
        cursor = conn.cursor()

        cursor.execute("SELECT * FROM grc_alert_rules ORDER BY created_at DESC")
        rules = [dict(row) for row in cursor.fetchall()]

        conn.close()
        return render_template('grc/alerts/rules.html', rules=rules)

    @app.route('/grc/alerts/events')
    @require_login
    @require_grc_permission('alert_events', 'view')
    def grc_alert_events_list():
        """List all alert events."""
        conn = get_grc_db()
        cursor = conn.cursor()

        cursor.execute("""
            SELECT ae.*, ar.rule_name
            FROM grc_alert_events ae
            JOIN grc_alert_rules ar ON ae.rule_id = ar.id
            ORDER BY ae.triggered_at DESC
        """)
        events = [dict(row) for row in cursor.fetchall()]

        conn.close()
        return render_template('grc/alerts/events.html', events=events)

    @app.route('/grc/high-risk')
    @require_login
    @require_grc_permission('high_risk_events', 'view')
    def grc_high_risk_list():
        """List all high-risk events."""
        conn = get_grc_db()
        cursor = conn.cursor()

        cursor.execute("""
            SELECT hre.*, u.full_name as user_name
            FROM grc_high_risk_events hre
            LEFT JOIN users u ON hre.user_id = u.id
            ORDER BY hre.detected_at DESC
        """)
        events = [dict(row) for row in cursor.fetchall()]

        conn.close()
        return render_template('grc/high_risk/index.html', events=events)

    @app.route('/grc/approval-matrix')
    @require_login
    @require_grc_permission('approval_matrix', 'view')
    def grc_approval_matrix_list():
        """List all approval matrix rules."""
        conn = get_grc_db()
        cursor = conn.cursor()

        cursor.execute("SELECT * FROM grc_approval_matrix_rules ORDER BY module, entity_type")
        rules = [dict(row) for row in cursor.fetchall()]

        conn.close()
        return render_template('grc/approval_matrix/index.html', rules=rules)

    @app.route('/grc/frameworks')
    @require_login
    @require_grc_permission('framework_templates', 'view')
    def grc_frameworks_list():
        """List all framework templates."""
        conn = get_grc_db()
        cursor = conn.cursor()

        cursor.execute("SELECT * FROM grc_framework_templates WHERE is_active = 1")
        frameworks = [dict(row) for row in cursor.fetchall()]

        conn.close()
        return render_template('grc/frameworks/index.html', frameworks=frameworks)

    @app.route('/grc/reports')
    @require_login
    @require_grc_permission('reports', 'view')
    def grc_reports_list():
        """GRC Reports center."""
        return render_template('grc/reports/index.html')

    @app.route('/grc/reports/risk-register')
    @require_login
    @require_grc_permission('reports', 'view')
    def grc_report_risk_register():
        """Risk register report."""
        conn = get_grc_db()
        cursor = conn.cursor()

        cursor.execute("""
            SELECT r.risk_code, r.title, r.risk_type, r.risk_level,
                   r.inherent_risk_score, r.residual_risk_score, r.status,
                   rc.category_name, u.full_name as owner_name
            FROM grc_risks r
            LEFT JOIN grc_risk_categories rc ON r.category_id = rc.id
            LEFT JOIN users u ON r.risk_owner_id = u.id
            WHERE r.status = 'active'
            ORDER BY r.inherent_risk_score DESC
        """)
        risks = [dict(row) for row in cursor.fetchall()]

        conn.close()

        columns = ['risk_code', 'title', 'risk_type', 'risk_level', 'inherent_risk_score',
                   'residual_risk_score', 'status', 'category_name', 'owner_name']

        if request.args.get('format') == 'excel_text':
            return export_to_excel_text(risks, columns, 'Risk_Register_Report')
        elif request.args.get('format') == 'excel_general':
            return export_to_excel_general(risks, columns, 'Risk_Register_Report')
        else:
            return render_template('grc/reports/risk_register.html', risks=risks)

    @app.route('/grc/reports/control-library')
    @require_login
    @require_grc_permission('reports', 'view')
    def grc_report_control_library():
        """Control library report."""
        conn = get_grc_db()
        cursor = conn.cursor()

        cursor.execute("""
            SELECT c.control_code, c.title, c.control_type, c.nature,
                   c.frequency, c.effectiveness_rating, c.status,
                   cc.category_name, u.full_name as owner_name
            FROM grc_controls c
            LEFT JOIN grc_control_categories cc ON c.category_id = cc.id
            LEFT JOIN users u ON c.owner_user_id = u.id
            WHERE c.is_active = 1
            ORDER BY c.control_code
        """)
        controls = [dict(row) for row in cursor.fetchall()]

        conn.close()

        columns = ['control_code', 'title', 'control_type', 'nature', 'frequency',
                   'effectiveness_rating', 'status', 'category_name', 'owner_name']

        if request.args.get('format') == 'excel_text':
            return export_to_excel_text(controls, columns, 'Control_Library_Report')
        elif request.args.get('format') == 'excel_general':
            return export_to_excel_general(controls, columns, 'Control_Library_Report')
        else:
            return render_template('grc/reports/control_library.html', controls=controls)

    @app.route('/grc/reports/sod-conflicts')
    @require_login
    @require_grc_permission('reports', 'view')
    def grc_report_sod_conflicts():
        """SoD conflicts report."""
        conn = get_grc_db()
        cursor = conn.cursor()

        cursor.execute("""
            SELECT sc.id, sr.rule_name, sc.status, sc.detected_at,
                   sc.resolution, sc.expiry_date, u.full_name as user_name
            FROM grc_sod_conflicts sc
            JOIN grc_sod_rules sr ON sc.rule_id = sr.id
            LEFT JOIN users u ON sc.user_id = u.id
            ORDER BY sc.detected_at DESC
        """)
        conflicts = [dict(row) for row in cursor.fetchall()]

        conn.close()

        columns = ['id', 'rule_name', 'status', 'detected_at', 'resolution', 'expiry_date', 'user_name']

        if request.args.get('format') == 'excel_text':
            return export_to_excel_text(conflicts, columns, 'SoD_Conflicts_Report')
        elif request.args.get('format') == 'excel_general':
            return export_to_excel_general(conflicts, columns, 'SoD_Conflicts_Report')
        else:
            return render_template('grc/reports/sod_conflicts.html', conflicts=conflicts)

    @app.route('/grc/export')
    @require_login
    @require_grc_permission('export', 'view')
    def grc_export_center():
        """GRC Export center."""
        return render_template('grc/export/index.html')

    @app.route('/grc/settings')
    @require_login
    @require_grc_permission('settings', 'view')
    def grc_settings():
        """GRC Settings page."""
        return render_template('grc/settings/index.html')

    @app.route('/grc/audit-logs')
    @require_login
    @require_grc_permission('audit_logs', 'view')
    def grc_audit_logs():
        """GRC Audit logs."""
        from database import get_all
        logs = get_all("""
            SELECT * FROM audit_log
            WHERE entity_type LIKE 'grc_%'
            ORDER BY created_at DESC
            LIMIT 100
        """)

        return render_template('grc/audit_logs/index.html', logs=logs)

    @app.route('/grc/flow')
    @require_login
    def grc_flow_notifications():
        """Flow notifications for GRC."""
        return render_template('grc/flow/index.html')

    print("[GRC Routes] Governance & Compliance routes registered successfully")
