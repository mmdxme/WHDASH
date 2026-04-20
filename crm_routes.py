"""
CRM Extended Routes Module
==========================
Flask route handlers for the enhanced CRM functionality.
Extends sales_routes.py with enterprise CRM features.

CRM Route Structure:
- /crm/dashboard/ - CRM Executive Dashboard
- /crm/leads/ - Lead Management
- /crm/opportunities/pipeline/ - Pipeline View
- /crm/activities/ - Activity Management
- /crm/complaints/ - Complaint Management
- /crm/key-accounts/ - Key Account Management
- /crm/forecasts/ - Forecasting
- /crm/journey/ - Customer Journey
- /crm/reports/ - CRM Reports
- /crm/settings/ - CRM Settings

Usage:
    from crm_routes import register_crm_routes
    register_crm_routes(app, require_login, require_permission, get_db)
"""

from flask import Flask, render_template, request, redirect, url_for, session, flash, jsonify, Response
from functools import wraps
from datetime import datetime, timedelta
import json
import csv
import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side


def register_crm_routes(app: Flask, require_login, require_permission, get_db):
    """Register all CRM management routes."""

    # =============================================================================
    # HELPER DECORATORS & FUNCTIONS
    # =============================================================================

    def crm_permission_required(resource, action='view'):
        """CRM-specific permission decorator."""
        def decorator(f):
            @wraps(f)
            def decorated_function(*args, **kwargs):
                user_id = session.get('user_id')
                if not user_id:
                    flash("Please login to access this page.", "error")
                    return redirect(url_for('login'))

                if not require_permission(user_id, 'crm', resource, action):
                    flash(f"Access Denied. You don't have permission to {action} {resource}.", "error")
                    return redirect(url_for('index'))

                return f(*args, **kwargs)
            return decorated_function
        return decorator

    def get_current_user():
        """Get current user info from session."""
        return {
            'id': session.get('user_id'),
            'username': session.get('username'),
            'role_id': session.get('role_id'),
            'role_name': session.get('role_name'),
            'company_id': session.get('company_id')
        }

    def get_translation(key, default=None):
        """Get translation for current language."""
        lang = session.get('language', 'en')
        from translations import get_translation as _get_trans
        return _get_trans(lang, key, default or key)

    # =============================================================================
    # CRM DASHBOARD
    # =============================================================================

    @app.route('/crm/')
    @app.route('/crm/dashboard/')
    @crm_permission_required('dashboard', 'view')
    def crm_dashboard():
        """CRM Executive Dashboard."""
        from crm_models import get_crm_dashboard_stats, get_crm_lookups

        user = get_current_user()
        date_from = request.args.get('date_from')
        date_to = request.args.get('date_to')

        # Default to last 30 days if no dates
        if not date_from:
            date_from = (datetime.now() - timedelta(days=30)).strftime('%Y-%m-%d')
        if not date_to:
            date_to = datetime.now().strftime('%Y-%m-%d')

        # Get stats
        stats = get_crm_dashboard_stats(
            salesperson_id=request.args.get('salesperson_id'),
            date_from=date_from,
            date_to=date_to
        )

        # Get lookups
        lookups = get_crm_lookups()

        # Get salespersons for filter
        with get_db() as db:
            salespersons = db.execute("""
                SELECT id, username FROM users
                WHERE role_id IN (SELECT id FROM roles WHERE role_name LIKE '%Sales%')
                ORDER BY username
            """).fetchall()

        return render_template('crm/dashboard.html',
                             stats=stats,
                             lookups=lookups,
                             salespersons=[dict(s) for s in salespersons],
                             date_from=date_from,
                             date_to=date_to,
                             user=user)

    # =============================================================================
    # CRM LEADS
    # =============================================================================

    @app.route('/crm/leads/')
    @crm_permission_required('leads', 'view')
    def crm_leads():
        """Lead list page."""
        from crm_models import get_leads, get_crm_lookups

        filters = {
            'search': request.args.get('search'),
            'status': request.args.get('status'),
            'source': request.args.get('source'),
            'priority': request.args.get('priority'),
            'score_min': request.args.get('score_min'),
            'score_max': request.args.get('score_max'),
            'assigned_to': request.args.get('assigned_to'),
            'date_from': request.args.get('date_from'),
            'date_to': request.args.get('date_to'),
            'is_hot': request.args.get('is_hot'),
        }

        page = int(request.args.get('page', 1))
        per_page = int(request.args.get('per_page', 50))

        result = get_leads(filters=filters, page=page, per_page=per_page)
        lookups = get_crm_lookups()

        with get_db() as db:
            salespersons = db.execute("""
                SELECT id, username FROM users
                WHERE role_id IN (SELECT id FROM roles WHERE role_name LIKE '%Sales%')
                ORDER BY username
            """).fetchall()

        return render_template('crm/leads/list.html',
                             leads=result['leads'],
                             total=result['total'],
                             page=result['page'],
                             pages=result['pages'],
                             filters=filters,
                             lookups=lookups,
                             salespersons=[dict(s) for s in salespersons])

    @app.route('/crm/leads/new/', methods=['GET', 'POST'])
    @crm_permission_required('leads', 'create')
    def crm_leads_new():
        """Create new lead."""
        from crm_models import create_lead, get_crm_lookups

        if request.method == 'POST':
            data = {
                'company_name': request.form.get('company_name'),
                'contact_name': request.form.get('contact_name'),
                'contact_title': request.form.get('contact_title'),
                'phone': request.form.get('phone'),
                'whatsapp': request.form.get('whatsapp'),
                'email': request.form.get('email'),
                'website': request.form.get('website'),
                'source': request.form.get('source'),
                'priority': request.form.get('priority', 'Medium'),
                'is_hot': request.form.get('is_hot'),
                'industry': request.form.get('industry'),
                'employee_count': request.form.get('employee_count'),
                'annual_revenue': request.form.get('annual_revenue'),
                'address': request.form.get('address'),
                'city': request.form.get('city'),
                'country': request.form.get('country'),
                'assigned_to': request.form.get('assigned_to') or session.get('user_id'),
                'estimated_value': request.form.get('estimated_value', 0),
                'currency': request.form.get('currency', 'AED'),
                'notes': request.form.get('notes'),
                'company_id': session.get('company_id')
            }

            lead_id = create_lead(data)
            flash("Lead created successfully!", "success")
            return redirect(url_for('crm_leads_view', lead_id=lead_id))

        lookups = get_crm_lookups()

        with get_db() as db:
            salespersons = db.execute("""
                SELECT id, username FROM users
                WHERE role_id IN (SELECT id FROM roles WHERE role_name LIKE '%Sales%')
                ORDER BY username
            """).fetchall()
            countries = db.execute("SELECT * FROM countries ORDER BY name").fetchall()

        return render_template('crm/leads/form.html',
                             lead=None,
                             lookups=lookups,
                             salespersons=[dict(s) for s in salespersons],
                             countries=[dict(c) for c in countries],
                             form_action='create')

    @app.route('/crm/leads/<int:lead_id>/')
    @crm_permission_required('leads', 'view')
    def crm_leads_view(lead_id):
        """View lead details."""
        from crm_models import get_lead_by_id, get_crm_lookups

        lead = get_lead_by_id(lead_id)
        if not lead:
            flash("Lead not found.", "error")
            return redirect(url_for('crm_leads'))

        lookups = get_crm_lookups()

        with get_db() as db:
            salespersons = db.execute("""
                SELECT id, username FROM users
                WHERE role_id IN (SELECT id FROM roles WHERE role_name LIKE '%Sales%')
                ORDER BY username
            """).fetchall()

        return render_template('crm/leads/view.html',
                             lead=lead,
                             lookups=lookups,
                             salespersons=[dict(s) for s in salespersons])

    @app.route('/crm/leads/<int:lead_id>/edit/', methods=['GET', 'POST'])
    @crm_permission_required('leads', 'edit')
    def crm_leads_edit(lead_id):
        """Edit lead."""
        from crm_models import get_lead_by_id, update_lead, get_crm_lookups

        lead = get_lead_by_id(lead_id)
        if not lead:
            flash("Lead not found.", "error")
            return redirect(url_for('crm_leads'))

        if request.method == 'POST':
            data = {
                'company_name': request.form.get('company_name'),
                'contact_name': request.form.get('contact_name'),
                'contact_title': request.form.get('contact_title'),
                'phone': request.form.get('phone'),
                'whatsapp': request.form.get('whatsapp'),
                'email': request.form.get('email'),
                'website': request.form.get('website'),
                'source': request.form.get('source'),
                'priority': request.form.get('priority'),
                'is_hot': request.form.get('is_hot'),
                'industry': request.form.get('industry'),
                'employee_count': request.form.get('employee_count'),
                'annual_revenue': request.form.get('annual_revenue'),
                'address': request.form.get('address'),
                'city': request.form.get('city'),
                'country': request.form.get('country'),
                'assigned_to': request.form.get('assigned_to'),
                'estimated_value': request.form.get('estimated_value', 0),
                'currency': request.form.get('currency', 'AED'),
                'notes': request.form.get('notes'),
            }

            update_lead(lead_id, data)
            flash("Lead updated successfully!", "success")
            return redirect(url_for('crm_leads_view', lead_id=lead_id))

        lookups = get_crm_lookups()

        with get_db() as db:
            salespersons = db.execute("""
                SELECT id, username FROM users
                WHERE role_id IN (SELECT id FROM roles WHERE role_name LIKE '%Sales%')
                ORDER BY username
            """).fetchall()
            countries = db.execute("SELECT * FROM countries ORDER BY name").fetchall()

        return render_template('crm/leads/form.html',
                             lead=lead,
                             lookups=lookups,
                             salespersons=[dict(s) for s in salespersons],
                             countries=[dict(c) for c in countries],
                             form_action='edit')

    @app.route('/crm/leads/<int:lead_id>/convert/', methods=['GET', 'POST'])
    @crm_permission_required('leads', 'edit')
    def crm_leads_convert(lead_id):
        """Convert lead to customer or opportunity."""
        from crm_models import get_lead_by_id, convert_lead_to_customer, convert_lead_to_opportunity

        lead = get_lead_by_id(lead_id)
        if not lead:
            flash("Lead not found.", "error")
            return redirect(url_for('crm_leads'))

        if request.method == 'POST':
            conversion_type = request.form.get('conversion_type')

            if conversion_type == 'customer':
                customer_id = convert_lead_to_customer(lead_id, {
                    'customer_type': request.form.get('customer_type', 'Retail'),
                    'notes': f"Converted from Lead {lead['lead_number']}"
                })
                flash(f"Lead converted to Customer successfully!", "success")
                return redirect(url_for('sales_customers_view', customer_id=customer_id))

            elif conversion_type == 'opportunity':
                opportunity_id = convert_lead_to_opportunity(lead_id, {
                    'estimated_value': request.form.get('estimated_value', lead.get('estimated_value', 0)),
                    'probability': request.form.get('probability', 20),
                    'expected_close_date': request.form.get('expected_close_date'),
                    'notes': f"Converted from Lead {lead['lead_number']}"
                })
                flash(f"Lead converted to Opportunity successfully!", "success")
                return redirect(url_for('sales_opportunities_view', opportunity_id=opportunity_id))

        return render_template('crm/leads/convert.html', lead=lead)

    @app.route('/crm/leads/<int:lead_id>/activity/', methods=['POST'])
    @crm_permission_required('leads', 'edit')
    def crm_leads_add_activity(lead_id):
        """Add activity to lead."""
        from crm_models import log_lead_activity

        activity_data = {
            'activity_type': request.form.get('activity_type'),
            'subject': request.form.get('subject'),
            'activity_date': request.form.get('activity_date', datetime.now().strftime('%Y-%m-%d')),
            'duration_minutes': request.form.get('duration_minutes'),
            'outcome': request.form.get('outcome'),
            'next_follow_up': request.form.get('next_follow_up'),
            'owner_id': session.get('user_id'),
            'notes': request.form.get('notes')
        }

        log_lead_activity(lead_id, activity_data)
        flash("Activity logged successfully!", "success")
        return redirect(url_for('crm_leads_view', lead_id=lead_id))

    # =============================================================================
    # CRM OPPORTUNITIES (PIPELINE)
    # =============================================================================

    @app.route('/crm/opportunities/')
    @crm_permission_required('opportunities', 'view')
    def crm_opportunities():
        """CRM opportunities list."""
        from crm_models import get_crm_opportunities, get_opportunity_pipeline_summary, get_crm_lookups

        filters = {
            'search': request.args.get('search'),
            'stage': request.args.get('stage'),
            'salesperson_id': request.args.get('salesperson_id'),
            'source': request.args.get('source'),
            'market': request.args.get('market'),
            'probability_min': request.args.get('probability_min'),
            'value_min': request.args.get('value_min'),
            'date_from': request.args.get('date_from'),
            'date_to': request.args.get('date_to'),
        }

        page = int(request.args.get('page', 1))
        per_page = int(request.args.get('per_page', 50))

        result = get_crm_opportunities(filters=filters, page=page, per_page=per_page)
        pipeline = get_opportunity_pipeline_summary()
        lookups = get_crm_lookups()

        with get_db() as db:
            salespersons = db.execute("""
                SELECT id, username FROM users
                WHERE role_id IN (SELECT id FROM roles WHERE role_name LIKE '%Sales%')
                ORDER BY username
            """).fetchall()

        return render_template('crm/opportunities/list.html',
                             opportunities=result['opportunities'],
                             total=result['total'],
                             page=result['page'],
                             pages=result['pages'],
                             filters=filters,
                             pipeline=pipeline,
                             lookups=lookups,
                             salespersons=[dict(s) for s in salespersons])

    @app.route('/crm/opportunities/pipeline/')
    @crm_permission_required('opportunities', 'view')
    def crm_pipeline():
        """Pipeline Kanban view."""
        from crm_models import get_crm_opportunities, get_opportunity_pipeline_summary, get_crm_lookups

        filters = {
            'stage_not': True  # Exclude closed deals for pipeline view
        }

        result = get_crm_opportunities(filters=filters, page=1, per_page=500)
        pipeline = get_opportunity_pipeline_summary()
        lookups = get_crm_lookups()

        # Group opportunities by stage
        stages = {}
        for stage in lookups['opportunity_stages']:
            if stage not in ['Won', 'Lost', 'Cancelled']:
                stages[stage] = [o for o in result['opportunities'] if o.get('stage') == stage]

        return render_template('crm/opportunities/pipeline.html',
                             stages=stages,
                             pipeline=pipeline,
                             lookups=lookups)

    # =============================================================================
    # CRM ACTIVITIES
    # =============================================================================

    @app.route('/crm/activities/')
    @crm_permission_required('activities', 'view')
    def crm_activities():
        """CRM activities list."""
        from crm_models import get_crm_activities, get_crm_lookups

        filters = {
            'search': request.args.get('search'),
            'activity_type': request.args.get('activity_type'),
            'owner_id': request.args.get('owner_id'),
            'customer_id': request.args.get('customer_id'),
            'date_from': request.args.get('date_from'),
            'date_to': request.args.get('date_to'),
        }

        page = int(request.args.get('page', 1))
        per_page = int(request.args.get('per_page', 50))

        result = get_crm_activities(filters=filters, page=page, per_page=per_page)
        lookups = get_crm_lookups()

        with get_db() as db:
            salespersons = db.execute("""
                SELECT id, username FROM users
                WHERE role_id IN (SELECT id FROM roles WHERE role_name LIKE '%Sales%')
                ORDER BY username
            """).fetchall()
            customers = db.execute("SELECT id, name FROM sales_customers WHERE is_active = 1 ORDER BY name").fetchall()

        return render_template('crm/activities/list.html',
                             activities=result['activities'],
                             total=result['total'],
                             page=result['page'],
                             pages=result['pages'],
                             filters=filters,
                             lookups=lookups,
                             salespersons=[dict(s) for s in salespersons],
                             customers=[dict(c) for c in customers])

    @app.route('/crm/activities/new/', methods=['GET', 'POST'])
    @crm_permission_required('activities', 'create')
    def crm_activities_new():
        """Create new activity."""
        from crm_models import create_crm_activity, get_crm_lookups

        if request.method == 'POST':
            activity_data = {
                'activity_type': request.form.get('activity_type'),
                'subject': request.form.get('subject'),
                'activity_date': request.form.get('activity_date', datetime.now().strftime('%Y-%m-%d')),
                'duration_minutes': request.form.get('duration_minutes'),
                'customer_id': request.form.get('customer_id'),
                'reference_type': request.form.get('reference_type'),
                'reference_id': request.form.get('reference_id'),
                'owner_id': session.get('user_id'),
                'outcome': request.form.get('outcome'),
                'next_follow_up': request.form.get('next_follow_up'),
                'notes': request.form.get('notes'),
                'company_id': session.get('company_id')
            }

            activity_id = create_crm_activity(activity_data)
            flash("Activity logged successfully!", "success")

            redirect_to = request.form.get('redirect_to')
            if redirect_to:
                return redirect(redirect_to)
            return redirect(url_for('crm_activities'))

        lookups = get_crm_lookups()

        with get_db() as db:
            customers = db.execute("SELECT id, name FROM sales_customers WHERE is_active = 1 ORDER BY name").fetchall()

        return render_template('crm/activities/form.html',
                             activity=None,
                             lookups=lookups,
                             customers=[dict(c) for c in customers],
                             redirect_to=request.args.get('redirect_to'))

    @app.route('/crm/activities/calendar/')
    @crm_permission_required('activities', 'view')
    def crm_activities_calendar():
        """Activities calendar view."""
        from crm_models import get_crm_activities, get_crm_lookups

        date_from = request.args.get('date_from', (datetime.now() - timedelta(days=7)).strftime('%Y-%m-%d'))
        date_to = request.args.get('date_to', (datetime.now() + timedelta(days=30)).strftime('%Y-%m-%d'))
        owner_id = request.args.get('owner_id')

        filters = {
            'date_from': date_from,
            'date_to': date_to,
            'owner_id': owner_id
        }

        result = get_crm_activities(filters=filters, page=1, per_page=500)
        lookups = get_crm_lookups()

        with get_db() as db:
            salespersons = db.execute("""
                SELECT id, username FROM users
                WHERE role_id IN (SELECT id FROM roles WHERE role_name LIKE '%Sales%')
                ORDER BY username
            """).fetchall()

        return render_template('crm/activities/calendar.html',
                             activities=result['activities'],
                             filters=filters,
                             lookups=lookups,
                             salespersons=[dict(s) for s in salespersons])

    # =============================================================================
    # CRM COMPLAINTS
    # =============================================================================

    @app.route('/crm/complaints/')
    @crm_permission_required('complaints', 'view')
    def crm_complaints():
        """Complaints list page."""
        from crm_models import get_complaints, get_crm_lookups

        filters = {
            'search': request.args.get('search'),
            'status': request.args.get('status'),
            'priority': request.args.get('priority'),
            'category': request.args.get('category'),
            'customer_id': request.args.get('customer_id'),
            'assigned_to': request.args.get('assigned_to'),
            'date_from': request.args.get('date_from'),
            'date_to': request.args.get('date_to'),
        }

        page = int(request.args.get('page', 1))
        per_page = int(request.args.get('per_page', 50))

        result = get_complaints(filters=filters, page=page, per_page=per_page)
        lookups = get_crm_lookups()

        with get_db() as db:
            salespersons = db.execute("""
                SELECT id, username FROM users
                WHERE role_id IN (SELECT id FROM roles WHERE role_name LIKE '%Sales%')
                ORDER BY username
            """).fetchall()
            customers = db.execute("SELECT id, name FROM sales_customers WHERE is_active = 1 ORDER BY name").fetchall()

        return render_template('crm/complaints/list.html',
                             complaints=result['complaints'],
                             total=result['total'],
                             page=result['page'],
                             pages=result['pages'],
                             filters=filters,
                             lookups=lookups,
                             salespersons=[dict(s) for s in salespersons],
                             customers=[dict(c) for c in customers])

    @app.route('/crm/complaints/new/', methods=['GET', 'POST'])
    @crm_permission_required('complaints', 'create')
    def crm_complaints_new():
        """Create new complaint."""
        from crm_models import create_complaint, get_crm_lookups

        if request.method == 'POST':
            data = {
                'customer_id': request.form.get('customer_id'),
                'order_id': request.form.get('order_id'),
                'delivery_id': request.form.get('delivery_id'),
                'category': request.form.get('category'),
                'priority': request.form.get('priority', 'Medium'),
                'subject': request.form.get('subject'),
                'description': request.form.get('description'),
                'assigned_to': request.form.get('assigned_to'),
                'resolution_target_date': request.form.get('resolution_target_date'),
                'company_id': session.get('company_id')
            }

            complaint_id = create_complaint(data)
            flash("Complaint created successfully!", "success")
            return redirect(url_for('crm_complaints_view', complaint_id=complaint_id))

        lookups = get_crm_lookups()

        with get_db() as db:
            salespersons = db.execute("""
                SELECT id, username FROM users
                WHERE role_id IN (SELECT id FROM roles WHERE role_name LIKE '%Sales%')
                ORDER BY username
            """).fetchall()
            customers = db.execute("SELECT id, name FROM sales_customers WHERE is_active = 1 ORDER BY name").fetchall()
            orders = db.execute("SELECT id, order_number, customer_id FROM sales_orders ORDER BY order_date DESC LIMIT 100").fetchall()

        return render_template('crm/complaints/form.html',
                             complaint=None,
                             lookups=lookups,
                             salespersons=[dict(s) for s in salespersons],
                             customers=[dict(c) for c in customers],
                             orders=[dict(o) for o in orders])

    @app.route('/crm/complaints/<int:complaint_id>/')
    @crm_permission_required('complaints', 'view')
    def crm_complaints_view(complaint_id):
        """View complaint details."""
        from crm_models import get_complaint_by_id, get_crm_lookups

        complaint = get_complaint_by_id(complaint_id)
        if not complaint:
            flash("Complaint not found.", "error")
            return redirect(url_for('crm_complaints'))

        lookups = get_crm_lookups()

        with get_db() as db:
            salespersons = db.execute("""
                SELECT id, username FROM users
                WHERE role_id IN (SELECT id FROM roles WHERE role_name LIKE '%Sales%')
                ORDER BY username
            """).fetchall()

        return render_template('crm/complaints/view.html',
                             complaint=complaint,
                             lookups=lookups,
                             salespersons=[dict(s) for s in salespersons])

    @app.route('/crm/complaints/<int:complaint_id>/update/', methods=['POST'])
    @crm_permission_required('complaints', 'edit')
    def crm_complaints_update(complaint_id):
        """Update complaint status."""
        from crm_models import update_complaint_status, add_complaint_response

        action = request.form.get('action')

        if action == 'status':
            update_complaint_status(
                complaint_id,
                request.form.get('status'),
                request.form.get('notes'),
                session.get('user_id')
            )
            flash("Complaint status updated!", "success")
        elif action == 'response':
            add_complaint_response(complaint_id, {
                'description': request.form.get('response_description'),
                'created_by': session.get('user_id')
            })
            flash("Response added!", "success")

        return redirect(url_for('crm_complaints_view', complaint_id=complaint_id))

    # =============================================================================
    # CRM KEY ACCOUNTS
    # =============================================================================

    @app.route('/crm/key-accounts/')
    @crm_permission_required('key_accounts', 'view')
    def crm_key_accounts():
        """Key accounts list."""
        from crm_models import get_key_accounts, get_crm_lookups

        filters = {
            'search': request.args.get('search'),
            'tier': request.args.get('tier'),
            'account_manager_id': request.args.get('account_manager_id'),
        }

        page = int(request.args.get('page', 1))
        per_page = int(request.args.get('per_page', 50))

        result = get_key_accounts(filters=filters, page=page, per_page=per_page)
        lookups = get_crm_lookups()

        with get_db() as db:
            account_managers = db.execute("""
                SELECT id, username FROM users
                WHERE role_id IN (SELECT id FROM roles WHERE role_name LIKE '%Sales%')
                ORDER BY username
            """).fetchall()

        return render_template('crm/key_accounts/list.html',
                             accounts=result['accounts'],
                             total=result['total'],
                             page=result['page'],
                             pages=result['pages'],
                             filters=filters,
                             lookups=lookups,
                             account_managers=[dict(am) for am in account_managers])

    @app.route('/crm/key-accounts/<int:account_id>/')
    @crm_permission_required('key_accounts', 'view')
    def crm_key_accounts_view(account_id):
        """View key account details."""
        from crm_models import get_account_plan, get_customer_360

        account = get_account_plan(account_id)
        if not account:
            flash("Key account not found.", "error")
            return redirect(url_for('crm_key_accounts'))

        # Get customer 360 if customer_id exists
        customer_360 = None
        if account.get('customer_id'):
            customer_360 = get_customer_360(account['customer_id'])

        return render_template('crm/key_accounts/view.html',
                             account=account,
                             customer_360=customer_360)

    # =============================================================================
    # CRM FORECASTS
    # =============================================================================

    @app.route('/crm/forecasts/')
    @crm_permission_required('forecasts', 'view')
    def crm_forecasts():
        """Forecasts view."""
        from crm_models import get_crm_forecasts, get_salesperson_forecast, get_crm_lookups

        filters = {
            'salesperson_id': request.args.get('salesperson_id'),
            'date_from': request.args.get('date_from', (datetime.now() - timedelta(days=365)).strftime('%Y-%m-%d')),
            'date_to': request.args.get('date_to', (datetime.now() + timedelta(days=180)).strftime('%Y-%m-%d')),
        }

        forecasts = get_crm_forecasts(filters)

        # If specific salesperson, get detailed forecast
        salesperson_forecast = None
        if filters.get('salesperson_id'):
            salesperson_forecast = get_salesperson_forecast(
                filters['salesperson_id'],
                datetime.now().year
            )

        lookups = get_crm_lookups()

        with get_db() as db:
            salespersons = db.execute("""
                SELECT id, username FROM users
                WHERE role_id IN (SELECT id FROM roles WHERE role_name LIKE '%Sales%')
                ORDER BY username
            """).fetchall()

        return render_template('crm/forecasts/list.html',
                             forecasts=forecasts,
                             salesperson_forecast=salesperson_forecast,
                             filters=filters,
                             lookups=lookups,
                             salespersons=[dict(s) for s in salespersons])

    # =============================================================================
    # CRM CUSTOMER JOURNEY
    # =============================================================================

    @app.route('/crm/journey/<int:customer_id>/')
    @crm_permission_required('journey', 'view')
    def crm_customer_journey(customer_id):
        """Customer journey view."""
        from crm_models import get_customer_journey, get_customer_360

        customer_360 = get_customer_360(customer_id)
        if not customer_360:
            flash("Customer not found.", "error")
            return redirect(url_for('sales_customers'))

        journey = get_customer_journey(customer_id)

        return render_template('crm/journey/view.html',
                             customer=customer_360,
                             journey=journey)

    # =============================================================================
    # CRM REPORTS
    # =============================================================================

    @app.route('/crm/reports/')
    @crm_permission_required('reports', 'view')
    def crm_reports():
        """CRM Reports center."""
        return render_template('crm/reports/index.html')

    @app.route('/crm/reports/lead-conversion/')
    @crm_permission_required('reports', 'view')
    def crm_reports_lead_conversion():
        """Lead conversion report."""
        from crm_models import get_lead_conversion_report

        filters = {
            'date_from': request.args.get('date_from', (datetime.now() - timedelta(days=365)).strftime('%Y-%m-%d')),
            'date_to': request.args.get('date_to', datetime.now().strftime('%Y-%m-%d')),
        }

        report = get_lead_conversion_report(filters)

        return render_template('crm/reports/lead_conversion.html',
                             report=report,
                             filters=filters)

    @app.route('/crm/reports/pipeline/')
    @crm_permission_required('reports', 'view')
    def crm_reports_pipeline():
        """Pipeline report."""
        from crm_models import get_sales_pipeline_report

        report = get_sales_pipeline_report()

        return render_template('crm/reports/pipeline.html', report=report)

    @app.route('/crm/reports/customer-analysis/')
    @crm_permission_required('reports', 'view')
    def crm_reports_customer_analysis():
        """Customer analysis report."""
        from crm_models import get_customer_analysis_report

        report = get_customer_analysis_report()

        return render_template('crm/reports/customer_analysis.html', report=report)

    # =============================================================================
    # CRM EXPORT ENDPOINTS
    # =============================================================================

    @app.route('/crm/export/leads/')
    @crm_permission_required('leads', 'view')
    def crm_export_leads():
        """Export leads to Excel."""
        from crm_models import get_exportable_leads

        filters = {
            'status': request.args.get('status'),
            'source': request.args.get('source'),
            'priority': request.args.get('priority'),
            'assigned_to': request.args.get('assigned_to'),
        }

        leads = get_exportable_leads(filters)

        wb = Workbook()
        ws = wb.active
        ws.title = "Leads"

        # Headers
        headers = ['Lead #', 'Company', 'Contact', 'Phone', 'Email', 'Source',
                   'Status', 'Priority', 'Score', 'Est. Value', 'Assigned To', 'Created']
        ws.append(headers)

        # Data
        for lead in leads:
            ws.append([
                lead.get('lead_number', ''),
                lead.get('company_name', ''),
                lead.get('contact_name', ''),
                lead.get('phone', ''),
                lead.get('email', ''),
                lead.get('source', ''),
                lead.get('status', ''),
                lead.get('priority', ''),
                lead.get('lead_score', ''),
                lead.get('estimated_value', ''),
                lead.get('assigned_to_name', ''),
                lead.get('created_at', '')[:10] if lead.get('created_at') else ''
            ])

        # Style header
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="CCCCCC")

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        return Response(
            output.getvalue(),
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            headers={'Content-Disposition': 'attachment; filename=leads_export.xlsx'}
        )

    @app.route('/crm/export/opportunities/')
    @crm_permission_required('opportunities', 'view')
    def crm_export_opportunities():
        """Export opportunities to Excel."""
        from crm_models import get_exportable_opportunities

        filters = {
            'stage': request.args.get('stage'),
            'salesperson_id': request.args.get('salesperson_id'),
        }

        opps = get_exportable_opportunities(filters)

        wb = Workbook()
        ws = wb.active
        ws.title = "Opportunities"

        headers = ['Opportunity #', 'Customer', 'Stage', 'Value', 'Probability',
                   'Weighted Value', 'Expected Close', 'Salesperson', 'Created']
        ws.append(headers)

        for opp in opps:
            ws.append([
                opp.get('opportunity_number', ''),
                opp.get('customer_name', ''),
                opp.get('stage', ''),
                opp.get('estimated_value', ''),
                opp.get('success_probability', ''),
                opp.get('weighted_value', ''),
                opp.get('expected_close_date', ''),
                opp.get('assigned_salesperson_name', ''),
                opp.get('created_at', '')[:10] if opp.get('created_at') else ''
            ])

        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="CCCCCC")

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        return Response(
            output.getvalue(),
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            headers={'Content-Disposition': 'attachment; filename=opportunities_export.xlsx'}
        )

    @app.route('/crm/export/activities/')
    @crm_permission_required('activities', 'view')
    def crm_export_activities():
        """Export activities to CSV."""
        from crm_models import get_exportable_activities

        filters = {
            'activity_type': request.args.get('activity_type'),
            'owner_id': request.args.get('owner_id'),
            'date_from': request.args.get('date_from'),
            'date_to': request.args.get('date_to'),
        }

        activities = get_exportable_activities(filters)

        output = io.StringIO()
        writer = csv.writer(output)

        writer.writerow(['Date', 'Type', 'Subject', 'Customer', 'Owner', 'Duration', 'Outcome', 'Notes'])

        for act in activities:
            writer.writerow([
                act.get('activity_date', ''),
                act.get('activity_type', ''),
                act.get('subject', ''),
                act.get('customer_name', ''),
                act.get('owner_name', ''),
                act.get('duration_minutes', ''),
                act.get('outcome', ''),
                act.get('notes', '')
            ])

        output.seek(0)

        return Response(
            output.getvalue(),
            mimetype='text/csv',
            headers={'Content-Disposition': 'attachment; filename=activities_export.csv'}
        )

    # =============================================================================
    # CRM CUSTOMER 360 VIEW
    # =============================================================================

    @app.route('/crm/customer-360/<int:customer_id>/')
    @crm_permission_required('customer_360', 'view')
    def crm_customer_360(customer_id):
        """Customer 360 view."""
        from crm_models import get_customer_360

        customer = get_customer_360(customer_id)
        if not customer:
            flash("Customer not found.", "error")
            return redirect(url_for('sales_customers'))

        return render_template('crm/customer_360/view.html', customer=customer)

    # =============================================================================
    # CRM SETTINGS
    # =============================================================================

    @app.route('/crm/settings/')
    @crm_permission_required('settings', 'view')
    def crm_settings():
        """CRM Settings page."""
        return render_template('crm/settings/index.html')

    @app.route('/crm/settings/sources/')
    @crm_permission_required('settings', 'edit')
    def crm_settings_sources():
        """Manage lead sources."""
        return render_template('crm/settings/sources.html')

    @app.route('/crm/settings/stages/')
    @crm_permission_required('settings', 'edit')
    def crm_settings_stages():
        """Manage opportunity stages."""
        return render_template('crm/settings/stages.html')

    # =============================================================================
    # CRM API ENDPOINTS
    # =============================================================================

    @app.route('/api/crm/lead/<int:lead_id>/score/')
    @crm_permission_required('leads', 'view')
    def api_crm_lead_score(lead_id):
        """Get lead score details."""
        from crm_models import get_lead_by_id

        lead = get_lead_by_id(lead_id)
        if not lead:
            return jsonify({'error': 'Lead not found'}), 404

        return jsonify({
            'lead_id': lead_id,
            'score': lead.get('lead_score', 0),
            'qualifications': lead.get('qualifications', []),
            'activities': len(lead.get('activities', []))
        })

    @app.route('/api/crm/customer/<int:customer_id>/engagement/')
    @crm_permission_required('customer_360', 'view')
    def api_crm_customer_engagement(customer_id):
        """Get customer engagement score."""
        from crm_models import calculate_customer_engagement_score

        engagement = calculate_customer_engagement_score(customer_id)
        return jsonify(engagement)

    @app.route('/api/crm/customer/<int:customer_id>/churn-risk/')
    @crm_permission_required('customer_360', 'view')
    def api_crm_customer_churn_risk(customer_id):
        """Get customer churn risk."""
        from crm_models import calculate_churn_risk

        risk = calculate_churn_risk(customer_id)
        return jsonify(risk)

    @app.route('/api/crm/pipeline/summary/')
    @crm_permission_required('opportunities', 'view')
    def api_crm_pipeline_summary():
        """Get pipeline summary for dashboard."""
        from crm_models import get_opportunity_pipeline_summary

        summary = get_opportunity_pipeline_summary()
        return jsonify(summary)
