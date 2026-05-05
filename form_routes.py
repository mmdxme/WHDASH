"""
Form Builder & Form Workflow Engine Routes
==========================================
Flask routes for the enterprise Form Builder system.
Includes template builder, form designer, submissions, approvals, and more.
"""

from flask import Blueprint, request, session, redirect, url_for, flash, render_template, jsonify, send_file
from functools import wraps
import json
import uuid
from datetime import datetime, timedelta
from werkzeug.utils import secure_filename
import os

# Import form models
from form_models import (
    initialize_form_builder_tables, get_form_template, get_form_templates,
    get_template_sections, get_section_fields, get_field_options, get_table_columns,
    get_field_rules, get_template_workflow, get_workflow_steps, get_step_transitions,
    get_submission, get_submission_values, get_submission_rows, get_row_values,
    get_submission_attachments, get_submission_signatures, get_submission_comments,
    get_submission_history, get_pending_approvals, get_user_drafts, get_user_submissions,
    get_submission_assignments, get_numbering_rule, generate_submission_number,
    create_form_template, update_form_template, create_form_section, create_form_field,
    create_field_option, create_table_column, create_field_rule, create_form_workflow,
    create_workflow_step, create_workflow_transition, create_form_submission,
    save_submission_value, add_submission_row, save_row_value, update_submission_status,
    assign_submission, complete_assignment, add_approval, add_signature, add_attachment,
    add_comment, create_form_notification, link_forms, log_form_access, create_numbering_rule,
    search_submissions, get_submission_stats, get_template_usage_stats,
    FIELD_TYPES, FIELD_TYPE_CATEGORIES, SUBMISSION_STATUSES, STATUS_COLORS,
    SUBMISSION_STATUS_DRAFT, SUBMISSION_STATUS_SUBMITTED, SUBMISSION_STATUS_APPROVED,
    SUBMISSION_STATUS_REJECTED, SUBMISSION_STATUS_RETURNED, SUBMISSION_STATUS_PENDING_SUPERVISOR,
    WORKFLOW_STEP_TYPES, WORKFLOW_ACTIONS, ROUTING_TYPES, SIGNATURE_TYPES,
    ATTACHMENT_TYPES, TEMPLATE_STATUS_DRAFT, TEMPLATE_STATUS_PUBLISHED, TEMPLATE_STATUS_ARCHIVED
)

from database import get_db_context, log_audit

# Import unified permission decorator
from permissions import require_permission
from export_utils import (
    send_export_response,
    get_export_columns
)


# =============================================================================
# EXPORT TYPES AND COLUMNS
# =============================================================================

FORM_EXPORT_TYPES = [
    'csv', 'excel_text', 'excel_general', 'json', 'xml', 'txt',
    'pdf', 'docx', 'html', 'printable', 'barcode', 'api',
    'email', 'zip', 'backup', 'sql_dump', 'dashboard',
    'summary', 'detailed', 'audit_log'
]

FORM_EXPORT_COLUMNS = {
    'templates': ['template_id', 'name', 'form_type', 'status', 'sections', 'fields', 'created_at'],
    'submissions': ['submission_id', 'template_name', 'status', 'submitter', 'submitted_at'],
    'drafts': ['draft_id', 'template_name', 'user', 'last_modified'],
    'approvals': ['approval_id', 'submission_id', 'step_name', 'approver', 'status', 'action_at']
}


# ============================================================================
# BLUEPRINT SETUP
# ============================================================================

form_bp = Blueprint('forms', __name__, url_prefix='/forms', template_folder='templates/forms')


def require_login(f):
    """Decorator to require login."""
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if 'user_id' not in session:
            flash("Please login to access this page.", "error")
            return redirect(url_for('login'))
        return f(*args, **kwargs)
    return decorated_function


def get_current_user():
    """Get current user from session."""
    if 'user_id' not in session:
        return None
    return {
        'id': session.get('user_id'),
        'username': session.get('username'),
        'role_id': session.get('role_id'),
        'company_id': session.get('company_id'),
        'department_id': session.get('department_id')
    }


# ============================================================================
# FORM TEMPLATE ROUTES
# ============================================================================

@form_bp.route('/')
@form_bp.route('/dashboard')
@require_login
def dashboard():
    """Form Builder Dashboard."""
    user = get_current_user()

    from form_models import get_all
    from database import get_db_context

    # Get stats
    stats = {
        'total_templates': 0,
        'active_templates': 0,
        'total_submissions': 0,
        'pending_submissions': 0,
        'pending_approvals': 0,
        'approved_count': 0,
        'rejected_count': 0,
        'approval_rate': 0
    }

    # Get template stats
    with get_db_context() as db:
        templates = db.execute("""
            SELECT status, COUNT(*) as count
            FROM form_templates
            GROUP BY status
        """).fetchall()
        for t in templates:
            stats['total_templates'] += t['count']
            if t['status'] == 'published':
                stats['active_templates'] = t['count']

        # Get submission stats
        submissions = db.execute("""
            SELECT status, COUNT(*) as count
            FROM form_submissions
            GROUP BY status
        """).fetchall()
        for s in submissions:
            stats['total_submissions'] += s['count']
            if s['status'] not in ['approved', 'rejected', 'cancelled', 'closed']:
                stats['pending_submissions'] += s['count']
            if s['status'] == 'approved':
                stats['approved_count'] = s['count']
            elif s['status'] == 'rejected':
                stats['rejected_count'] = s['count']

        # Get pending approvals count for current user
        pending = db.execute("""
            SELECT COUNT(*) as cnt FROM form_assignments
            WHERE assigned_to = ? AND is_active = 1 AND status = 'pending'
        """, (user['id'],)).fetchone()
        stats['pending_approvals'] = pending['cnt'] if pending else 0

        # Calculate approval rate
        total_decisions = stats['approved_count'] + stats['rejected_count']
        if total_decisions > 0:
            stats['approval_rate'] = round((stats['approved_count'] / total_decisions) * 100)

    # Get recent submissions
    recent_submissions = get_all("""
        SELECT fs.*, ft.form_title
        FROM form_submissions fs
        JOIN form_templates ft ON fs.template_id = ft.id
        ORDER BY fs.created_at DESC
        LIMIT 10
    """)

    # Get my pending tasks
    my_pending = get_all("""
        SELECT fs.*, ft.form_title,
               CASE WHEN fs.due_date < date('now') THEN 1 ELSE 0 END as is_overdue
        FROM form_submissions fs
        JOIN form_templates ft ON fs.template_id = ft.id
        JOIN form_assignments fa ON fs.id = fa.submission_id
        WHERE fa.assigned_to = ? AND fa.is_active = 1 AND fa.status = 'pending'
        ORDER BY fs.priority DESC, fs.due_date ASC
        LIMIT 10
    """, (user['id'],))

    # Get template categories
    template_categories = get_all("""
        SELECT COALESCE(category, 'Uncategorized') as category, COUNT(*) as count
        FROM form_templates
        WHERE status = 'published'
        GROUP BY category
        ORDER BY count DESC
    """)

    # Get status breakdown
    status_breakdown = get_all("""
        SELECT status, COUNT(*) as count
        FROM form_submissions
        GROUP BY status
        ORDER BY count DESC
    """)

    # Get recent activity
    recent_activity = get_all("""
        SELECT fa.action, fa.created_at, u.username
        FROM form_audit fa
        LEFT JOIN users u ON fa.user_id = u.id
        ORDER BY fa.created_at DESC
        LIMIT 10
    """)

    return render_template('forms/form_dashboard.html',
                          stats=stats,
                          recent_submissions=recent_submissions,
                          my_pending=my_pending,
                          template_categories=template_categories,
                          status_breakdown=status_breakdown,
                          recent_activity=recent_activity)


# ============================================================================
# MY FORMS ROUTES
# ============================================================================

@form_bp.route('/my/drafts')
@require_login
def my_drafts():
    """Show user's draft submissions."""
    user = get_current_user()
    drafts = get_user_drafts(user['id'])
    return render_template('forms/my_drafts.html', submissions=drafts)


@form_bp.route('/my/submissions')
@require_login
def my_submissions():
    """Show user's submission history."""
    user = get_current_user()
    submissions = get_user_submissions(user['id'])
    return render_template('forms/my_submissions.html', submissions=submissions)


@form_bp.route('/my/pending')
@require_login
def my_pending_list():
    """Show user's pending approvals."""
    user = get_current_user()
    pending = get_pending_approvals(user['id'])

    # Calculate counts for stats
    from database import get_db_context
    today = datetime.now().strftime('%Y-%m-%d')

    overdue_count = 0
    due_today_count = 0
    due_week_count = len(pending)

    with get_db_context() as db:
        for p in pending:
            if p.get('due_date'):
                if p['due_date'] < today:
                    overdue_count += 1
                elif p['due_date'] == today:
                    due_today_count += 1

    return render_template('forms/my_pending.html',
                          submissions=pending,
                          overdue_count=overdue_count,
                          due_today_count=due_today_count,
                          due_week_count=due_week_count)


@form_bp.route('/templates')
@require_login
def templates_list():
    """List all form templates."""
    user = get_current_user()
    status_filter = request.args.get('status')
    department_filter = request.args.get('department')
    category_filter = request.args.get('category')
    
    templates = get_form_templates(
        status=status_filter,
        department=department_filter,
        category=category_filter
    )
    
    return render_template('forms/templates_list.html',
                          templates=templates,
                          status_filter=status_filter,
                          department_filter=department_filter,
                          category_filter=category_filter,
                          template_statuses=[TEMPLATE_STATUS_DRAFT, TEMPLATE_STATUS_PUBLISHED, TEMPLATE_STATUS_ARCHIVED])


@form_bp.route('/templates/new', methods=['GET', 'POST'])
@require_login
def templates_new():
    """Create a new form template."""
    user = get_current_user()
    
    if request.method == 'POST':
        data = request.form.to_dict()
        
        # Generate template code if not provided
        if not data.get('template_code'):
            data['template_code'] = f"TPL-{datetime.now().strftime('%Y%m%d%H%M%S')}"
        
        data['created_by'] = user['id']
        
        # Handle metadata
        metadata = {
            'allow_edit_after_submit': data.get('allow_edit_after_submit', '0') == '1',
            'allow_delete_draft': data.get('allow_delete_draft', '0') == '1',
            'show_watermark': data.get('show_watermark', '0') == '1',
            'require_correction_reason': data.get('require_correction_reason', '0') == '1',
        }
        data['metadata'] = json.dumps(metadata)
        
        template_id = create_form_template(data)
        
        # Create default numbering rule if prefix provided
        if data.get('numbering_prefix'):
            create_numbering_rule(
                template_id,
                prefix=data.get('numbering_prefix'),
                suffix=data.get('numbering_suffix', ''),
                date_format=data.get('numbering_date_format', 'YYYY'),
                include_year=1,
                include_month=0,
                include_day=0,
                padding_length=int(data.get('padding_length', 5))
            )
        
        log_audit('form_template', template_id, 'CREATE', user['id'],
                  notes=f"Created template: {data.get('form_title')}")
        
        flash("Form template created successfully.", "success")
        return redirect(url_for('forms.templates_edit', template_id=template_id))
    
    return render_template('forms/template_new.html')


@form_bp.route('/templates/<int:template_id>')
@require_login
def templates_view(template_id):
    """View a form template."""
    template = get_form_template(template_id)
    if not template:
        flash("Template not found.", "error")
        return redirect(url_for('forms.templates_list'))
    
    sections = get_template_sections(template_id)
    
    # Get fields for each section
    for section in sections:
        section['fields'] = get_section_fields(section['id'])
        for field in section['fields']:
            if field['field_type'] in ['dropdown', 'radio', 'multi_select']:
                field['options'] = get_field_options(field['id'])
            if field['field_type'] == 'table':
                field['columns'] = get_table_columns(field['id'])
    
    return render_template('forms/template_view.html',
                          template=template,
                          sections=sections)


@form_bp.route('/templates/<int:template_id>/edit', methods=['GET', 'POST'])
@require_login
def templates_edit(template_id):
    """Edit a form template."""
    user = get_current_user()
    template = get_form_template(template_id)
    if not template:
        flash("Template not found.", "error")
        return redirect(url_for('forms.templates_list'))
    
    sections = get_template_sections(template_id)
    
    # Get fields for each section
    for section in sections:
        section['fields'] = get_section_fields(section['id'])
        for field in section['fields']:
            if field['field_type'] in ['dropdown', 'radio', 'multi_select']:
                field['options'] = get_field_options(field['id'])
            if field['field_type'] == 'table':
                field['columns'] = get_table_columns(field['id'])
            field['rules'] = get_field_rules(field['id'])
    
    if request.method == 'POST':
        data = request.form.to_dict()
        
        # Handle metadata
        metadata = {
            'allow_edit_after_submit': data.get('allow_edit_after_submit', '0') == '1',
            'allow_delete_draft': data.get('allow_delete_draft', '0') == '1',
            'show_watermark': data.get('show_watermark', '0') == '1',
            'require_correction_reason': data.get('require_correction_reason', '0') == '1',
        }
        data['metadata'] = json.dumps(metadata)
        
        update_form_template(template_id, data)
        
        log_audit('form_template', template_id, 'UPDATE', user['id'],
                  notes=f"Updated template: {template['form_title']}")
        
        flash("Template updated successfully.", "success")
        return redirect(url_for('forms.templates_edit', template_id=template_id))
    
    return render_template('forms/template_edit.html',
                          template=template,
                          sections=sections,
                          field_types=FIELD_TYPES,
                          field_type_categories=FIELD_TYPE_CATEGORIES,
                          workflow_step_types=WORKFLOW_STEP_TYPES,
                          routing_types=ROUTING_TYPES)


@form_bp.route('/templates/<int:template_id>/publish', methods=['POST'])
@require_login
def templates_publish(template_id):
    """Publish a form template."""
    user = get_current_user()
    template = get_form_template(template_id)
    if not template:
        return jsonify({'success': False, 'message': 'Template not found'})
    
    update_form_template(template_id, {
        'status': TEMPLATE_STATUS_PUBLISHED,
        'published_at': datetime.now().isoformat()
    })
    
    log_audit('form_template', template_id, 'PUBLISH', user['id'],
              notes=f"Published template: {template['form_title']}")
    
    flash("Template published successfully.", "success")
    return redirect(url_for('forms.templates_view', template_id=template_id))


@form_bp.route('/templates/<int:template_id>/archive', methods=['POST'])
@require_login
def templates_archive(template_id):
    """Archive a form template."""
    user = get_current_user()
    template = get_form_template(template_id)
    if not template:
        return jsonify({'success': False, 'message': 'Template not found'})
    
    update_form_template(template_id, {
        'status': TEMPLATE_STATUS_ARCHIVED,
        'is_archived': 1,
        'archived_at': datetime.now().isoformat()
    })
    
    log_audit('form_template', template_id, 'ARCHIVE', user['id'],
              notes=f"Archived template: {template['form_title']}")
    
    flash("Template archived successfully.", "success")
    return redirect(url_for('forms.templates_list'))


@form_bp.route('/templates/<int:template_id>/duplicate', methods=['POST'])
@require_login
def templates_duplicate(template_id):
    """Duplicate a form template."""
    user = get_current_user()
    template = get_form_template(template_id)
    if not template:
        return jsonify({'success': False, 'message': 'Template not found'})
    
    # Create new template with duplicated data
    new_code = f"{template['template_code']}-COPY-{datetime.now().strftime('%Y%m%d%H%M%S')}"
    new_title = f"{template['form_title']} (Copy)"
    
    data = template.copy()
    data['template_code'] = new_code
    data['form_title'] = new_title
    data['status'] = TEMPLATE_STATUS_DRAFT
    data['created_by'] = user['id']
    data['created_at'] = datetime.now().isoformat()
    data['version'] = 1
    data.pop('id', None)
    data.pop('published_at', None)
    data.pop('archived_at', None)
    
    new_id = create_form_template(data)
    
    # Note: Would need to also duplicate sections, fields, workflows in a full implementation
    
    log_audit('form_template', new_id, 'DUPLICATE', user['id'],
              notes=f"Duplicated from template {template_id}: {template['form_title']}")
    
    flash("Template duplicated successfully.", "success")
    return redirect(url_for('forms.templates_edit', template_id=new_id))


# ============================================================================
# SECTION ROUTES
# ============================================================================

@form_bp.route('/templates/<int:template_id>/sections/add', methods=['POST'])
@require_login
def sections_add(template_id):
    """Add a new section to a template."""
    user = get_current_user()
    data = request.get_json()
    
    # Get next section order
    sections = get_template_sections(template_id)
    next_order = len(sections) if sections else 0
    
    section_data = {
        'template_id': template_id,
        'section_code': data.get('section_code', f"SEC-{datetime.now().strftime('%Y%m%d%H%M%S')}"),
        'title': data.get('title', 'New Section'),
        'title_ar': data.get('title_ar'),
        'description': data.get('description'),
        'section_order': next_order,
        'is_collapsible': data.get('is_collapsible', 1),
        'is_repeatable': data.get('is_repeatable', 0),
        'max_repeat_count': data.get('max_repeat_count'),
        'layout_type': data.get('layout_type', 'vertical'),
        'column_count': data.get('column_count', 1),
        'is_visible': data.get('is_visible', 1),
        'is_required': data.get('is_required', 1)
    }
    
    section_id = create_form_section(section_data)
    
    log_audit('form_section', section_id, 'CREATE', user['id'],
              notes=f"Added section to template {template_id}")
    
    return jsonify({'success': True, 'section_id': section_id})


@form_bp.route('/sections/<int:section_id>', methods=['GET'])
@require_login
def sections_get(section_id):
    """Get section details."""
    from database import get_one
    
    section = get_one("SELECT * FROM form_sections WHERE id = ?", (section_id,))
    if section:
        section['fields'] = get_section_fields(section_id)
        for field in section['fields']:
            if field['field_type'] in ['dropdown', 'radio', 'multi_select']:
                field['options'] = get_field_options(field['id'])
            if field['field_type'] == 'table':
                field['columns'] = get_table_columns(field['id'])
    
    return jsonify({'success': True, 'section': section})


@form_bp.route('/sections/<int:section_id>', methods=['PUT'])
@require_login
def sections_update(section_id):
    """Update a section."""
    user = get_current_user()
    data = request.get_json()
    
    from database import get_db_context
    
    fields = []
    values = []
    for key in ['title', 'title_ar', 'description', 'section_order', 'is_collapsible',
                'is_repeatable', 'max_repeat_count', 'layout_type', 'column_count',
                'is_visible', 'is_required']:
        if key in data:
            fields.append(f"{key} = ?")
            values.append(data[key])
    
    if fields:
        fields.append("updated_at = ?")
        values.append(datetime.now().isoformat())
        values.append(section_id)
        
        with get_db_context() as db:
            db.execute(f"UPDATE form_sections SET {', '.join(fields)} WHERE id = ?", values)
            db.commit()
    
    log_audit('form_section', section_id, 'UPDATE', user['id'])
    
    return jsonify({'success': True})


@form_bp.route('/sections/<int:section_id>', methods=['DELETE'])
@require_login
def sections_delete(section_id):
    """Delete a section."""
    user = get_current_user()
    
    from database import get_db_context
    with get_db_context() as db:
        db.execute("DELETE FROM form_sections WHERE id = ?", (section_id,))
        db.commit()
    
    log_audit('form_section', section_id, 'DELETE', user['id'])
    
    return jsonify({'success': True})


@form_bp.route('/sections/reorder', methods=['POST'])
@require_login
def sections_reorder():
    """Reorder sections in a template."""
    user = get_current_user()
    data = request.get_json()
    
    section_orders = data.get('section_orders', [])
    
    from database import get_db_context
    with get_db_context() as db:
        for section_id, new_order in section_orders:
            db.execute("UPDATE form_sections SET section_order = ? WHERE id = ?", (new_order, section_id))
        db.commit()
    
    return jsonify({'success': True})


# ============================================================================
# FIELD ROUTES
# ============================================================================

@form_bp.route('/sections/<int:section_id>/fields/add', methods=['POST'])
@require_login
def fields_add(section_id):
    """Add a new field to a section."""
    user = get_current_user()
    data = request.get_json()
    
    # Get next field order
    fields = get_section_fields(section_id)
    next_order = len(fields) if fields else 0
    
    # Generate field code if not provided
    field_code = data.get('field_code')
    if not field_code:
        base_code = data.get('label', 'field').lower().replace(' ', '_')
        field_code = f"fld_{base_code}_{datetime.now().strftime('%Y%m%d%H%M%S')}"
    
    field_data = {
        'section_id': section_id,
        'field_code': field_code,
        'field_type': data.get('field_type', 'text'),
        'label': data.get('label', 'New Field'),
        'label_ar': data.get('label_ar'),
        'placeholder': data.get('placeholder'),
        'help_text': data.get('help_text'),
        'default_value': data.get('default_value'),
        'validation_rules': data.get('validation_rules'),
        'field_order': next_order,
        'field_width': data.get('field_width', 12),
        'is_required': data.get('is_required', 0),
        'is_readonly': data.get('is_readonly', 0),
        'is_hidden': data.get('is_hidden', 0),
        'is_unique': data.get('is_unique', 0),
        'min_value': data.get('min_value'),
        'max_value': data.get('max_value'),
        'min_length': data.get('min_length'),
        'max_length': data.get('max_length'),
        'accepted_file_types': data.get('accepted_file_types'),
        'max_file_size_mb': data.get('max_file_size_mb'),
        'data_source': data.get('data_source'),
        'data_source_config': data.get('data_source_config'),
        'calculation_formula': data.get('calculation_formula'),
        'conditional_logic': data.get('conditional_logic')
    }
    
    field_id = create_form_field(field_data)
    
    # Add options if provided (for dropdown, radio, multi_select)
    if data.get('options'):
        for idx, opt in enumerate(data['options']):
            create_field_option({
                'field_id': field_id,
                'option_value': opt.get('value', str(idx)),
                'option_label': opt.get('label', f"Option {idx + 1}"),
                'option_label_ar': opt.get('label_ar'),
                'option_order': idx,
                'is_default': 1 if idx == 0 else 0
            })
    
    # Add table columns if this is a table field
    if data.get('field_type') == 'table' and data.get('columns'):
        for idx, col in enumerate(data['columns']):
            create_table_column({
                'field_id': field_id,
                'column_code': col.get('code', f"col_{idx}"),
                'column_label': col.get('label', f"Column {idx + 1}"),
                'column_label_ar': col.get('label_ar'),
                'column_type': col.get('type', 'text'),
                'column_order': idx,
                'column_width': col.get('width'),
                'is_required': col.get('required', 0),
                'is_readonly': col.get('readonly', 0),
                'default_value': col.get('default'),
                'min_value': col.get('min_value'),
                'max_value': col.get('max_value'),
                'data_source': col.get('data_source'),
                'calculation_formula': col.get('formula')
            })
    
    log_audit('form_field', field_id, 'CREATE', user['id'],
              notes=f"Added field to section {section_id}")
    
    return jsonify({'success': True, 'field_id': field_id})


@form_bp.route('/fields/<int:field_id>', methods=['GET'])
@require_login
def fields_get(field_id):
    """Get field details."""
    from database import get_one
    
    field = get_one("SELECT * FROM form_fields WHERE id = ?", (field_id,))
    if field:
        if field['field_type'] in ['dropdown', 'radio', 'multi_select']:
            field['options'] = get_field_options(field_id)
        if field['field_type'] == 'table':
            field['columns'] = get_table_columns(field_id)
        field['rules'] = get_field_rules(field_id)
    
    return jsonify({'success': True, 'field': field})


@form_bp.route('/fields/<int:field_id>', methods=['PUT'])
@require_login
def fields_update(field_id):
    """Update a field."""
    user = get_current_user()
    data = request.get_json()
    
    from database import get_db_context
    import json
    
    fields = []
    values = []
    for key in ['label', 'label_ar', 'placeholder', 'help_text', 'default_value',
                'validation_rules', 'field_order', 'field_width', 'is_required',
                'is_readonly', 'is_hidden', 'is_unique', 'min_value', 'max_value',
                'min_length', 'max_length', 'accepted_file_types', 'max_file_size_mb',
                'data_source', 'calculation_formula']:
        if key in data:
            fields.append(f"{key} = ?")
            values.append(data[key])
    
    if 'data_source_config' in data:
        fields.append("data_source_config = ?")
        values.append(json.dumps(data['data_source_config']))
    
    if 'conditional_logic' in data:
        fields.append("conditional_logic = ?")
        values.append(json.dumps(data['conditional_logic']))
    
    if fields:
        fields.append("updated_at = ?")
        values.append(datetime.now().isoformat())
        values.append(field_id)
        
        with get_db_context() as db:
            db.execute(f"UPDATE form_fields SET {', '.join(fields)} WHERE id = ?", values)
            db.commit()
    
    log_audit('form_field', field_id, 'UPDATE', user['id'])
    
    return jsonify({'success': True})


@form_bp.route('/fields/<int:field_id>', methods=['DELETE'])
@require_login
def fields_delete(field_id):
    """Delete a field."""
    user = get_current_user()
    
    from database import get_db_context
    with get_db_context() as db:
        db.execute("DELETE FROM form_fields WHERE id = ?", (field_id,))
        db.commit()
    
    log_audit('form_field', field_id, 'DELETE', user['id'])
    
    return jsonify({'success': True})


@form_bp.route('/fields/<int:field_id>/options', methods=['POST'])
@require_login
def fields_add_options(field_id):
    """Add options to a field."""
    user = get_current_user()
    data = request.get_json()
    
    options = data.get('options', [])
    # Clear existing options
    from database import get_db_context
    with get_db_context() as db:
        db.execute("DELETE FROM field_options WHERE field_id = ?", (field_id,))
        for idx, opt in enumerate(options):
            db.execute("""
                INSERT INTO field_options (field_id, option_value, option_label, option_label_ar, option_order, is_default)
                VALUES (?, ?, ?, ?, ?, ?)
            """, (field_id, opt.get('value', str(idx)), opt.get('label'), opt.get('label_ar'), idx, 1 if opt.get('default') else 0))
        db.commit()
    
    return jsonify({'success': True})


# ============================================================================
# WORKFLOW ROUTES
# ============================================================================

@form_bp.route('/templates/<int:template_id>/workflow', methods=['GET', 'POST'])
@require_login
def workflow_edit(template_id):
    """Edit workflow for a template."""
    user = get_current_user()
    template = get_form_template(template_id)
    if not template:
        flash("Template not found.", "error")
        return redirect(url_for('forms.templates_list'))
    
    workflow = get_template_workflow(template_id)
    steps = []
    if workflow:
        steps = get_workflow_steps(workflow['id'])
    
    if request.method == 'POST':
        data = request.form.to_dict()
        
        # Create or update workflow
        if workflow:
            workflow_id = workflow['id']
        else:
            workflow_id = create_form_workflow({
                'template_id': template_id,
                'workflow_name': data.get('workflow_name', f"Workflow for {template['form_title']}"),
                'workflow_name_ar': data.get('workflow_name_ar'),
                'description': data.get('description'),
                'workflow_type': data.get('workflow_type', 'serial'),
                'is_default': 1,
                'requires_all_approvals': data.get('requires_all_approvals', '0') == '1',
                'allow_partial_approval': data.get('allow_partial_approval', '0') == '1',
                'return_to_originator': data.get('return_to_originator', '0') == '1'
            })
        
        flash("Workflow saved successfully.", "success")
        return redirect(url_for('forms.workflow_edit', template_id=template_id))
    
    return render_template('forms/workflow_edit.html',
                          template=template,
                          workflow=workflow,
                          steps=steps,
                          workflow_step_types=WORKFLOW_STEP_TYPES,
                          workflow_actions=WORKFLOW_ACTIONS,
                          routing_types=ROUTING_TYPES)


@form_bp.route('/workflows/<int:workflow_id>/steps/add', methods=['POST'])
@require_login
def workflow_steps_add(workflow_id):
    """Add a step to a workflow."""
    user = get_current_user()
    data = request.get_json()
    
    # Get next step order
    steps = get_workflow_steps(workflow_id)
    next_order = len(steps) if steps else 0
    
    step_data = {
        'workflow_id': workflow_id,
        'step_order': next_order,
        'step_name': data.get('step_name', f"Step {next_order + 1}"),
        'step_name_ar': data.get('step_name_ar'),
        'step_type': data.get('step_type', 'approval'),
        'approver_type': data.get('approver_type', 'user'),
        'approver_id': data.get('approver_id'),
        'approver_role_id': data.get('approver_role_id'),
        'approver_department_id': data.get('approver_department_id'),
        'routing_condition': data.get('routing_condition'),
        'is_parallel_step': data.get('is_parallel_step', 0),
        'parallel_approvers': data.get('parallel_approvers'),
        'sla_hours': data.get('sla_hours'),
        'escalation_step_id': data.get('escalation_step_id')
    }
    
    step_id = create_workflow_step(step_data)
    
    log_audit('workflow_step', step_id, 'CREATE', user['id'],
              notes=f"Added step to workflow {workflow_id}")
    
    return jsonify({'success': True, 'step_id': step_id})


@form_bp.route('/workflows/<int:workflow_id>/steps', methods=['GET'])
@require_login
def workflow_steps_list(workflow_id):
    """Get all steps for a workflow."""
    steps = get_workflow_steps(workflow_id)
    return jsonify({'success': True, 'steps': steps})


@form_bp.route('/steps/<int:step_id>', methods=['PUT'])
@require_login
def workflow_step_update(step_id):
    """Update a workflow step."""
    user = get_current_user()
    data = request.get_json()
    
    from database import get_db_context
    import json
    
    fields = []
    values = []
    for key in ['step_name', 'step_name_ar', 'step_type', 'approver_type',
                'approver_id', 'approver_role_id', 'approver_department_id',
                'routing_condition', 'is_parallel_step', 'sla_hours', 'step_status']:
        if key in data:
            fields.append(f"{key} = ?")
            values.append(data[key])
    
    if 'parallel_approvers' in data:
        fields.append("parallel_approvers = ?")
        values.append(json.dumps(data['parallel_approvers']))
    
    if fields:
        values.append(step_id)
        with get_db_context() as db:
            db.execute(f"UPDATE workflow_steps SET {', '.join(fields)} WHERE id = ?", values)
            db.commit()
    
    log_audit('workflow_step', step_id, 'UPDATE', user['id'])
    
    return jsonify({'success': True})


@form_bp.route('/steps/<int:step_id>', methods=['DELETE'])
@require_login
def workflow_step_delete(step_id):
    """Delete a workflow step."""
    user = get_current_user()
    
    from database import get_db_context
    with get_db_context() as db:
        db.execute("DELETE FROM workflow_steps WHERE id = ?", (step_id,))
        db.commit()
    
    log_audit('workflow_step', step_id, 'DELETE', user['id'])
    
    return jsonify({'success': True})


# ============================================================================
# FORM SUBMISSION ROUTES
# ============================================================================

@form_bp.route('/submissions')
@require_login
def submissions_list():
    """List form submissions."""
    user = get_current_user()
    
    # Get filters
    status_filter = request.args.get('status')
    template_filter = request.args.get('template')
    search = request.args.get('search')
    my_pending = request.args.get('my_pending') == '1'
    
    # Get user's pending approvals
    if my_pending:
        pending = get_pending_approvals(user['id'])
        return render_template('forms/submissions_list.html',
                              submissions=pending,
                              my_pending=True,
                              status_filter='pending')
    
    submissions = search_submissions(
        template_id=template_filter,
        status=status_filter,
        search_text=search,
        submitted_by=user['id'] if request.args.get('my_submissions') == '1' else None,
        assignee_id=user['id'] if request.args.get('my_assignments') == '1' else None
    )
    
    return render_template('forms/submissions_list.html',
                          submissions=submissions,
                          status_filter=status_filter,
                          template_filter=template_filter,
                          search=search)


@form_bp.route('/submissions/new/<int:template_id>', methods=['GET', 'POST'])
@require_login
def submissions_new(template_id):
    """Create a new form submission."""
    user = get_current_user()
    template = get_form_template(template_id)
    if not template:
        flash("Template not found.", "error")
        return redirect(url_for('forms.templates_list'))
    
    if template['status'] != TEMPLATE_STATUS_PUBLISHED:
        flash("This template is not available for submissions.", "warning")
        return redirect(url_for('forms.templates_list'))
    
    sections = get_template_sections(template_id)
    for section in sections:
        section['fields'] = get_section_fields(section['id'])
        for field in section['fields']:
            if field['field_type'] in ['dropdown', 'radio', 'multi_select']:
                field['options'] = get_field_options(field['id'])
            if field['field_type'] == 'table':
                field['columns'] = get_table_columns(field['id'])
    
    workflow = get_template_workflow(template_id)
    
    if request.method == 'POST':
        data = request.form.to_dict()
        
        # Create submission
        submission_data = {
            'template_id': template_id,
            'title': data.get('title', template['form_title']),
            'status': SUBMISSION_STATUS_DRAFT,
            'submitted_by': user['id'],
            'workflow_id': workflow['id'] if workflow else None,
            'department_id': user['department_id'],
            'priority': data.get('priority', 'medium')
        }
        
        submission_id = create_form_submission(submission_data)
        
        # Save field values
        for section in sections:
            for field in section['fields']:
                field_key = f"field_{field['id']}"
                if field_key in data:
                    save_submission_value(
                        submission_id, field['id'], field['field_code'],
                        field['field_type'], data[field_key], user['id']
                    )
        
        # Save table rows
        for section in sections:
            for field in section['fields']:
                if field['field_type'] == 'table':
                    row_count = int(request.form.get(f"row_count_{field['id']}", 0))
                    columns = get_table_columns(field['id'])
                    
                    for row_idx in range(row_count):
                        row_id = add_submission_row(submission_id, field['id'], row_idx)
                        
                        for col in columns:
                            col_key = f"table_{field['id']}_{row_idx}_{col['id']}"
                            if col_key in data:
                                save_row_value(row_id, col['id'], col['column_code'], data[col_key])
        
        # Handle attachments
        attachments = request.files.getlist('attachments')
        for attachment in attachments:
            if attachment.filename:
                # Save file and create attachment record
                filename = secure_filename(attachment.filename)
                upload_dir = os.path.join('static', 'uploads', 'forms', str(submission_id))
                os.makedirs(upload_dir, exist_ok=True)
                filepath = os.path.join(upload_dir, filename)
                attachment.save(filepath)
                
                add_attachment(
                    submission_id=submission_id,
                    file_name=filename,
                    file_path=filepath,
                    file_type=attachment.content_type,
                    file_size=attachment.content_length,
                    uploaded_by=user['id'],
                    attachment_type='general'
                )
        
        log_audit('form_submission', submission_id, 'CREATE', user['id'],
                  notes=f"Created submission for template {template_id}")
        
        flash("Form saved as draft.", "success")
        return redirect(url_for('forms.submissions_view', submission_id=submission_id))
    
    return render_template('forms/submission_new.html',
                          template=template,
                          sections=sections,
                          workflow=workflow)


@form_bp.route('/submissions/<int:submission_id>')
@require_login
def submissions_view(submission_id):
    """View a form submission."""
    user = get_current_user()

    submission = get_submission(submission_id)
    if not submission:
        flash("Submission not found.", "error")
        return redirect(url_for('forms.submissions_list'))

    template = get_form_template(submission['template_id'])
    sections = get_template_sections(submission['template_id'])

    # Get field values
    values = get_submission_values(submission_id)
    values_dict = {v['field_id']: v['field_value'] for v in values}

    # Get sections with fields
    table_data = {}
    for section in sections:
        section['fields'] = get_section_fields(section['id'])
        for field in section['fields']:
            field['current_value'] = values_dict.get(field['id'])
            if field['field_type'] in ['dropdown', 'radio', 'multi_select']:
                field['options'] = get_field_options(field['id'])
            if field['field_type'] == 'table':
                field['columns'] = get_table_columns(field['id'])
                rows = get_submission_rows(submission_id, field['id'])
                table_data[field['id']] = rows

                # Convert rows to list format for table display
                field['rows'] = []
                for row in rows:
                    row_values = get_row_values(row['id'])
                    row_dict = {rv['column_code']: rv['field_value'] for rv in row_values}
                    field['rows'].append(row_dict)

    # Get attachments, signatures, comments, history
    attachments = get_submission_attachments(submission_id)
    signatures = get_submission_signatures(submission_id)
    comments = get_submission_comments(submission_id)
    history = get_submission_history(submission_id)
    assignments = get_submission_assignments(submission_id)

    # Get current assignment
    current_assignment = None
    for a in assignments:
        if a.get('is_active') == 1 or a.get('status') == 'pending':
            current_assignment = a
            # Get assigned user name
            from database import get_one
            assigned_user = get_one("SELECT username FROM users WHERE id = ?", (a['assigned_to'],))
            if assigned_user:
                current_assignment['assigned_to_name'] = assigned_user['username']
            break

    # Get workflow progress
    workflow_progress = {'completed_steps': [], 'pending_step': None, 'total_steps': 0, 'current_step': 0}
    workflow = get_template_workflow(submission['template_id'])
    if workflow:
        steps = get_workflow_steps(workflow['id'])
        workflow_progress['total_steps'] = len(steps)
        workflow_progress['current_step'] = submission.get('current_step_id') or 0

        for step in steps:
            step_approvals = [h for h in history if h.get('workflow_step_id') == step['id']]
            if step_approvals:
                workflow_progress['completed_steps'].append({
                    'step': step,
                    'approvals': step_approvals
                })
            elif workflow_progress['pending_step'] is None:
                workflow_progress['pending_step'] = step

    # Get related forms
    from database import get_all
    related_forms = get_all("""
        SELECT * FROM form_submissions
        WHERE parent_submission_id = ? OR related_form_id = ?
        LIMIT 5
    """, (submission_id, submission_id))

    # Get activity log
    activity_log = get_all("""
        SELECT fa.*, u.username
        FROM form_audit fa
        LEFT JOIN users u ON fa.user_id = u.id
        WHERE fa.entity_type = 'form_submission' AND fa.entity_id = ?
        ORDER BY fa.created_at DESC
        LIMIT 20
    """, (submission_id,))

    # Check permissions
    can_edit = submission['status'] == 'draft' and submission['submitted_by'] == user['id']
    can_submit = submission['status'] == 'draft' and submission['submitted_by'] == user['id']
    can_approve = False
    if current_assignment and current_assignment['assigned_to'] == user['id']:
        can_approve = True

    # Log access
    log_form_access(submission_id, user['id'], 'view')

    return render_template('forms/submission_view.html',
                          submission=submission,
                          template=template,
                          sections=sections,
                          values=values_dict,
                          table_data=table_data,
                          attachments=attachments,
                          signatures=signatures,
                          comments=comments,
                          history=history,
                          assignments=assignments,
                          current_assignment=current_assignment,
                          workflow_progress=workflow_progress,
                          related_forms=related_forms,
                          activity_log=activity_log,
                          can_edit=can_edit,
                          can_submit=can_submit,
                          can_approve=can_approve,
                          status_colors=STATUS_COLORS)


@form_bp.route('/submissions/<int:submission_id>/edit', methods=['GET', 'POST'])
@require_login
def submissions_edit(submission_id):
    """Edit a form submission (draft only)."""
    user = get_current_user()
    
    submission = get_submission(submission_id)
    if not submission:
        flash("Submission not found.", "error")
        return redirect(url_for('forms.submissions_list'))
    
    if submission['status'] != SUBMISSION_STATUS_DRAFT:
        flash("Only draft submissions can be edited.", "warning")
        return redirect(url_for('forms.submissions_view', submission_id=submission_id))
    
    if submission['submitted_by'] != user['id']:
        flash("You can only edit your own submissions.", "error")
        return redirect(url_for('forms.submissions_view', submission_id=submission_id))
    
    template = get_form_template(submission['template_id'])
    sections = get_template_sections(submission['template_id'])
    
    # Get field values
    values = get_submission_values(submission_id)
    values_dict = {v['field_id']: v['field_value'] for v in values}
    
    for section in sections:
        section['fields'] = get_section_fields(section['id'])
        for field in section['fields']:
            field['current_value'] = values_dict.get(field['id'])
            if field['field_type'] in ['dropdown', 'radio', 'multi_select']:
                field['options'] = get_field_options(field['id'])
            if field['field_type'] == 'table':
                field['columns'] = get_table_columns(field['id'])
                field['rows'] = get_submission_rows(submission_id, field['id'])
    
    if request.method == 'POST':
        # Update field values
        for section in sections:
            for field in section['fields']:
                field_key = f"field_{field['id']}"
                if field_key in request.form:
                    save_submission_value(
                        submission_id, field['id'], field['field_code'],
                        field['field_type'], request.form[field_key], user['id']
                    )
        
        flash("Form updated successfully.", "success")
        return redirect(url_for('forms.submissions_view', submission_id=submission_id))
    
    return render_template('forms/submission_edit.html',
                          submission=submission,
                          template=template,
                          sections=sections)


@form_bp.route('/submissions/<int:submission_id>/submit', methods=['POST'])
@require_login
def submissions_submit(submission_id):
    """Submit a form for approval."""
    user = get_current_user()
    
    submission = get_submission(submission_id)
    if not submission:
        return jsonify({'success': False, 'message': 'Submission not found'})
    
    # Update status
    update_submission_status(submission_id, SUBMISSION_STATUS_SUBMITTED, user['id'],
                            reason='Form submitted for approval')
    
    # Get workflow and assign to first approver
    workflow = get_template_workflow(submission['template_id'])
    if workflow:
        steps = get_workflow_steps(workflow['id'])
        if steps:
            first_step = steps[0]
            if first_step['approver_type'] == 'user' and first_step['approver_id']:
                assign_submission(
                    submission_id=submission_id,
                    assigned_to=first_step['approver_id'],
                    assigned_by=user['id'],
                    workflow_step_id=first_step['id']
                )
            elif first_step['approver_type'] == 'role' and first_step['approver_role_id']:
                # Would need to find users with this role and assign to them
                pass
    
    # Create notification
    create_form_notification(
        submission_id=submission_id,
        notification_type='submission',
        title='Form Submitted',
        message=f"Your form {submission['submission_number']} has been submitted.",
        recipient_id=user['id']
    )
    
    log_audit('form_submission', submission_id, 'SUBMIT', user['id'])
    
    flash("Form submitted successfully.", "success")
    return redirect(url_for('forms.submissions_view', submission_id=submission_id))


@form_bp.route('/submissions/<int:submission_id>/approve', methods=['POST'])
@require_login
def submissions_approve(submission_id):
    """Approve a form submission."""
    user = get_current_user()
    comment = request.form.get('comment', '')
    
    submission = get_submission(submission_id)
    if not submission:
        return jsonify({'success': False, 'message': 'Submission not found'})
    
    # Record approval
    workflow_step_id = submission.get('current_step_id')
    add_approval(submission_id, workflow_step_id or 0, user['id'], 'approve', comment)
    
    # Update status
    update_submission_status(submission_id, SUBMISSION_STATUS_APPROVED, user['id'],
                            reason=comment)
    
    # Create notification
    create_form_notification(
        submission_id=submission_id,
        notification_type='approval',
        title='Form Approved',
        message=f"Your form {submission['submission_number']} has been approved.",
        recipient_id=submission['submitted_by']
    )
    
    log_audit('form_submission', submission_id, 'APPROVE', user['id'], notes=comment)
    
    flash("Form approved successfully.", "success")
    return redirect(url_for('forms.submissions_view', submission_id=submission_id))


@form_bp.route('/submissions/<int:submission_id>/reject', methods=['POST'])
@require_login
def submissions_reject(submission_id):
    """Reject a form submission."""
    user = get_current_user()
    comment = request.form.get('comment', '')
    reason = request.form.get('reason', '')
    
    if not comment:
        flash("A reason is required when rejecting.", "error")
        return redirect(url_for('forms.submissions_view', submission_id=submission_id))
    
    submission = get_submission(submission_id)
    if not submission:
        return jsonify({'success': False, 'message': 'Submission not found'})
    
    # Record rejection
    workflow_step_id = submission.get('current_step_id')
    add_approval(submission_id, workflow_step_id or 0, user['id'], 'reject', comment)
    
    # Update status
    update_submission_status(submission_id, SUBMISSION_STATUS_REJECTED, user['id'],
                            reason=reason or comment)
    
    # Create notification
    create_form_notification(
        submission_id=submission_id,
        notification_type='rejection',
        title='Form Rejected',
        message=f"Your form {submission['submission_number']} has been rejected. Reason: {comment}",
        recipient_id=submission['submitted_by']
    )
    
    log_audit('form_submission', submission_id, 'REJECT', user['id'], notes=comment)
    
    flash("Form rejected.", "success")
    return redirect(url_for('forms.submissions_view', submission_id=submission_id))


@form_bp.route('/submissions/<int:submission_id>/return', methods=['POST'])
@require_login
def submissions_return(submission_id):
    """Return a form for correction."""
    user = get_current_user()
    comment = request.form.get('comment', '')
    reason = request.form.get('reason', '')
    
    if not comment:
        flash("A reason is required when returning a form.", "error")
        return redirect(url_for('forms.submissions_view', submission_id=submission_id))
    
    submission = get_submission(submission_id)
    if not submission:
        return jsonify({'success': False, 'message': 'Submission not found'})
    
    # Update status
    update_submission_status(submission_id, SUBMISSION_STATUS_RETURNED, user['id'],
                            reason=reason or comment)
    
    # Create notification
    create_form_notification(
        submission_id=submission_id,
        notification_type='return',
        title='Form Returned for Correction',
        message=f"Your form {submission['submission_number']} has been returned. Reason: {comment}",
        recipient_id=submission['submitted_by']
    )
    
    log_audit('form_submission', submission_id, 'RETURN', user['id'], notes=comment)
    
    flash("Form returned for correction.", "success")
    return redirect(url_for('forms.submissions_view', submission_id=submission_id))


@form_bp.route('/submissions/<int:submission_id>/cancel', methods=['POST'])
@require_login
def submissions_cancel(submission_id):
    """Cancel a form submission."""
    user = get_current_user()
    reason = request.form.get('reason', '')
    
    submission = get_submission(submission_id)
    if not submission:
        return jsonify({'success': False, 'message': 'Submission not found'})
    
    if submission['submitted_by'] != user['id']:
        return jsonify({'success': False, 'message': 'You can only cancel your own submissions'})
    
    # Update status
    from database import get_db_context
    with get_db_context() as db:
        db.execute("""
            UPDATE form_submissions SET status = 'cancelled', cancelled_at = ?,
            cancelled_by = ?, cancellation_reason = ?, updated_at = ?
            WHERE id = ?
        """, (datetime.now().isoformat(), user['id'], reason, datetime.now().isoformat(), submission_id))
        db.execute("""
            INSERT INTO form_status_history (submission_id, from_status, to_status, changed_by, reason)
            VALUES (?, 'draft', 'cancelled', ?, ?)
        """, (submission_id, user['id'], reason))
        db.commit()
    
    log_audit('form_submission', submission_id, 'CANCEL', user['id'], notes=reason)
    
    flash("Form cancelled.", "success")
    return redirect(url_for('forms.submissions_list'))


# ============================================================================
# SIGNATURE ROUTES
# ============================================================================

@form_bp.route('/submissions/<int:submission_id>/sign', methods=['POST'])
@require_login
def submissions_sign(submission_id):
    """Add a signature to a submission."""
    user = get_current_user()
    data = request.get_json()
    
    signature_type = data.get('signature_type', 'typed')
    signature_data = data.get('signature_data')
    signature_image_path = data.get('signature_image_path')
    
    # Get user info for signature
    from database import get_one
    user_data = get_one("SELECT username FROM users WHERE id = ?", (user['id'],))
    signed_by_name = user_data['username'] if user_data else user['username']
    
    signature_id = add_signature(
        submission_id=submission_id,
        signature_type=signature_type,
        signed_by=user['id'],
        signed_by_name=signed_by_name,
        signature_data=signature_data,
        signature_image_path=signature_image_path,
        workflow_step_id=data.get('workflow_step_id'),
        ip_address=request.remote_addr,
        user_agent=request.headers.get('User-Agent')
    )
    
    log_audit('form_signature', signature_id, 'SIGN', user['id'],
              notes=f"Signature added to submission {submission_id}")
    
    return jsonify({'success': True, 'signature_id': signature_id})


# ============================================================================
# COMMENT ROUTES
# ============================================================================

@form_bp.route('/submissions/<int:submission_id>/comments/add', methods=['POST'])
@require_login
def submissions_add_comment(submission_id):
    """Add a comment to a submission."""
    user = get_current_user()
    data = request.get_json()
    
    comment = add_comment(
        submission_id=submission_id,
        comment_text=data.get('comment_text', ''),
        created_by=user['id'],
        comment_type=data.get('comment_type', 'general'),
        is_internal=data.get('is_internal', 0),
        parent_comment_id=data.get('parent_comment_id'),
        mentioned_users=data.get('mentioned_users')
    )
    
    log_audit('form_comment', comment, 'CREATE', user['id'])
    
    return jsonify({'success': True, 'comment_id': comment})


# ============================================================================
# ATTACHMENT ROUTES
# ============================================================================

@form_bp.route('/submissions/<int:submission_id>/attachments/add', methods=['POST'])
@require_login
def submissions_add_attachment(submission_id):
    """Add an attachment to a submission."""
    user = get_current_user()
    
    if 'file' not in request.files:
        return jsonify({'success': False, 'message': 'No file provided'})
    
    file = request.files['file']
    if file.filename == '':
        return jsonify({'success': False, 'message': 'No file selected'})
    
    # Save file
    filename = secure_filename(file.filename)
    upload_dir = os.path.join('static', 'uploads', 'forms', str(submission_id))
    os.makedirs(upload_dir, exist_ok=True)
    filepath = os.path.join(upload_dir, filename)
    file.save(filepath)
    
    # Get file size
    file.seek(0, 2)
    file_size = file.tell()
    file.seek(0)
    
    attachment_id = add_attachment(
        submission_id=submission_id,
        file_name=filename,
        file_path=filepath,
        file_type=file.content_type,
        file_size=file_size,
        uploaded_by=user['id'],
        section_id=request.form.get('section_id'),
        field_id=request.form.get('field_id'),
        row_id=request.form.get('row_id'),
        attachment_type=request.form.get('attachment_type', 'general'),
        is_required=request.form.get('is_required', 0)
    )
    
    log_audit('form_attachment', attachment_id, 'UPLOAD', user['id'],
              notes=f"Attachment added to submission {submission_id}: {filename}")
    
    return jsonify({'success': True, 'attachment_id': attachment_id})


@form_bp.route('/attachments/<int:attachment_id>/download')
@require_login
def attachments_download(attachment_id):
    """Download an attachment."""
    from database import get_one
    
    attachment = get_one("SELECT * FROM form_submission_attachments WHERE id = ?", (attachment_id,))
    if not attachment:
        flash("Attachment not found.", "error")
        return redirect(url_for('forms.submissions_list'))
    
    return send_file(attachment['file_path'], as_attachment=True, download_name=attachment['file_name'])


@form_bp.route('/attachments/<int:attachment_id>', methods=['DELETE'])
@require_login
def attachments_delete(attachment_id):
    """Delete an attachment."""
    user = get_current_user()
    
    from database import get_db_context
    with get_db_context() as db:
        db.execute("DELETE FROM form_submission_attachments WHERE id = ?", (attachment_id,))
        db.commit()
    
    log_audit('form_attachment', attachment_id, 'DELETE', user['id'])
    
    return jsonify({'success': True})


# ============================================================================
# DASHBOARD ROUTE
# ============================================================================

@form_bp.route('/reports')
@require_login
def reports():
    """Form Builder reports dashboard."""
    user = get_current_user()

    from form_models import get_all

    # Get comprehensive stats
    with get_db_context() as db:
        # Submission by status
        status_stats = db.execute("""
            SELECT status, COUNT(*) as count
            FROM form_submissions
            GROUP BY status
        """).fetchall()

        # Submissions over time (last 30 days)
        time_stats = db.execute("""
            SELECT DATE(created_at) as date, COUNT(*) as count
            FROM form_submissions
            WHERE created_at >= date('now', '-30 days')
            GROUP BY DATE(created_at)
            ORDER BY date
        """).fetchall()

        # Top templates by usage
        top_templates = db.execute("""
            SELECT ft.form_title, COUNT(*) as submission_count
            FROM form_submissions fs
            JOIN form_templates ft ON fs.template_id = ft.id
            GROUP BY ft.id
            ORDER BY submission_count DESC
            LIMIT 10
        """).fetchall()

        # Department breakdown
        dept_stats = db.execute("""
            SELECT COALESCE(d.name, 'Unassigned') as department, COUNT(*) as count
            FROM form_submissions fs
            LEFT JOIN departments d ON fs.department_id = d.id
            GROUP BY d.id
            ORDER BY count DESC
        """).fetchall()

        # Average approval time
        avg_approval_time = db.execute("""
            SELECT AVG(
                (julianday(MAX(created_at)) - julianday(MIN(created_at))) * 24
            ) as avg_hours
            FROM form_approvals
            WHERE submission_id IN (
                SELECT id FROM form_submissions WHERE status = 'approved'
            )
        """).fetchone()

        # Overdue rate
        overdue_count = db.execute("""
            SELECT COUNT(*) as cnt
            FROM form_submissions
            WHERE due_date < date('now')
            AND status NOT IN ('approved', 'rejected', 'cancelled', 'closed')
        """).fetchone()

        template_stats = get_template_usage_stats()
        submission_stats = get_submission_stats()

    return render_template('forms/reports.html',
                          template_stats=template_stats,
                          submission_stats=submission_stats,
                          status_stats=status_stats,
                          time_stats=time_stats,
                          top_templates=top_templates,
                          dept_stats=dept_stats,
                          avg_approval_time=avg_approval_time['avg_hours'] if avg_approval_time and avg_approval_time['avg_hours'] else 0,
                          overdue_count=overdue_count['cnt'] if overdue_count else 0,
                          status_colors=STATUS_COLORS)


@form_bp.route('/reports/search')
@require_login
def reports_search():
    """Search submissions with filters."""
    user = get_current_user()

    results = search_submissions(
        template_id=request.args.get('template_id'),
        status=request.args.get('status'),
        department_id=request.args.get('department_id'),
        submitted_by=request.args.get('submitted_by'),
        date_from=request.args.get('date_from'),
        date_to=request.args.get('date_to'),
        search_text=request.args.get('search'),
        priority=request.args.get('priority'),
        overdue=request.args.get('overdue') == '1'
    )

    return render_template('forms/reports_search.html',
                          results=results,
                          status_colors=STATUS_COLORS)


# ============================================================================
# EXPORT ROUTES
# ============================================================================

@form_bp.route('/submissions/<int:submission_id>/export/pdf')
@require_login
def export_submission_pdf(submission_id):
    """Export submission as PDF."""
    submission = get_submission(submission_id)
    if not submission:
        flash("Submission not found.", "error")
        return redirect(url_for('forms.submissions_list'))

    template = get_form_template(submission['template_id'])
    sections = get_template_sections(submission['template_id'])

    # Get values
    values = get_submission_values(submission_id)
    values_dict = {v['field_id']: v['field_value'] for v in values}

    for section in sections:
        section['fields'] = get_section_fields(section['id'])
        for field in section['fields']:
            field['current_value'] = values_dict.get(field['id'])

    # Generate HTML for PDF (could use weasyprint or similar)
    html = render_template('forms/export/pdf_template.html',
                          submission=submission,
                          template=template,
                          sections=sections,
                          values=values_dict)

    # For now, return HTML (in production, convert to PDF)
    from flask import make_response
    response = make_response(html)
    response.headers['Content-Type'] = 'text/html'
    return response


@form_bp.route('/submissions/<int:submission_id>/export/excel')
@require_login
def export_submission_excel(submission_id):
    """Export submission as Excel."""
    submission = get_submission(submission_id)
    if not submission:
        flash("Submission not found.", "error")
        return redirect(url_for('forms.submissions_list'))

    from openpyxl import Workbook
    from io import BytesIO

    wb = Workbook()
    ws = wb.active
    ws.title = "Submission"

    # Add header info
    ws['A1'] = 'Submission Number'
    ws['B1'] = submission['submission_number']
    ws['A2'] = 'Form'
    ws['B2'] = submission.get('title', 'N/A')
    ws['A3'] = 'Status'
    ws['B3'] = submission['status']
    ws['A4'] = 'Date'
    ws['B4'] = submission['created_at'][:10] if submission.get('created_at') else ''

    # Add field values
    values = get_submission_values(submission_id)
    row = 6
    for v in values:
        ws[f'A{row}'] = v['field_code']
        ws[f'B{row}'] = v['field_value']
        row += 1

    # Save to BytesIO
    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(output,
                    mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                    as_attachment=True,
                    download_name=f"submission_{submission['submission_number']}.xlsx")


@form_bp.route('/reports/export')
@require_login
def reports_export():
    """Export report data as CSV."""
    from io import StringIO
    import csv

    # Get filter parameters
    template_id = request.args.get('template_id')
    status = request.args.get('status')
    date_from = request.args.get('date_from')
    date_to = request.args.get('date_to')

    # Search submissions
    results = search_submissions(
        template_id=template_id,
        status=status,
        date_from=date_from,
        date_to=date_to,
        limit=10000
    )

    # Create CSV
    output = StringIO()
    writer = csv.writer(output)

    # Write header
    writer.writerow([
        'Submission Number', 'Form Title', 'Status', 'Priority',
        'Submitted By', 'Department', 'Created At', 'Submitted At',
        'Due Date', 'Current Step'
    ])

    # Write data
    for sub in results:
        writer.writerow([
            sub.get('submission_number', ''),
            sub.get('form_title', ''),
            sub.get('status', ''),
            sub.get('priority', ''),
            sub.get('submitted_by_name', ''),
            sub.get('department_name', ''),
            sub.get('created_at', '')[:16] if sub.get('created_at') else '',
            sub.get('submitted_at', '')[:16] if sub.get('submitted_at') else '',
            sub.get('due_date', ''),
            sub.get('current_step_name', '')
        ])

    output.seek(0)

    return send_file(output,
                    mimetype='text/csv',
                    as_attachment=True,
                    download_name=f"forms_report_{datetime.now().strftime('%Y%m%d')}.csv")


# ============================================================================
# API ROUTES (JSON)
# ============================================================================

@form_bp.route('/api/templates')
@require_login
def api_templates_list():
    """API: List all templates."""
    templates = get_form_templates(status='published')
    return jsonify({'success': True, 'templates': templates})


@form_bp.route('/api/templates/<int:template_id>')
@require_login
def api_template_detail(template_id):
    """API: Get template with sections and fields."""
    template = get_form_template(template_id)
    if not template:
        return jsonify({'success': False, 'message': 'Template not found'}), 404

    sections = get_template_sections(template_id)
    for section in sections:
        section['fields'] = get_section_fields(section['id'])
        for field in section['fields']:
            if field['field_type'] in ['dropdown', 'radio', 'multi_select']:
                field['options'] = get_field_options(field['id'])
            if field['field_type'] == 'table':
                field['columns'] = get_table_columns(field['id'])

    workflow = get_template_workflow(template_id)

    return jsonify({
        'success': True,
        'template': template,
        'sections': sections,
        'workflow': workflow
    })


@form_bp.route('/api/submissions/<int:submission_id>/values')
@require_login
def api_submission_values(submission_id):
    """API: Get all values for a submission."""
    submission = get_submission(submission_id)
    if not submission:
        return jsonify({'success': False, 'message': 'Submission not found'}), 404

    values = get_submission_values(submission_id)
    values_dict = {v['field_id']: v['field_value'] for v in values}

    # Also get table rows
    sections = get_template_sections(submission['template_id'])
    table_data = {}
    for section in sections:
        for field in get_section_fields(section['id']):
            if field['field_type'] == 'table':
                rows = get_submission_rows(submission_id, field['id'])
                table_data[field['id']] = rows

    return jsonify({
        'success': True,
        'submission': submission,
        'values': values_dict,
        'table_data': table_data
    })


@form_bp.route('/api/submissions/<int:submission_id>/submit', methods=['POST'])
@require_login
def api_submit_submission(submission_id):
    """API: Submit a form for approval."""
    user = get_current_user()
    submission = get_submission(submission_id)

    if not submission:
        return jsonify({'success': False, 'message': 'Submission not found'}), 404

    if submission['status'] != SUBMISSION_STATUS_DRAFT:
        return jsonify({'success': False, 'message': 'Only drafts can be submitted'}), 400

    # Update status
    update_submission_status(submission_id, SUBMISSION_STATUS_SUBMITTED, user['id'],
                            reason='Form submitted for approval')

    # Get workflow and assign to first approver
    workflow = get_template_workflow(submission['template_id'])
    if workflow:
        steps = get_workflow_steps(workflow['id'])
        if steps:
            first_step = steps[0]
            if first_step['approver_type'] == 'user' and first_step['approver_id']:
                assign_submission(
                    submission_id=submission_id,
                    assigned_to=first_step['approver_id'],
                    assigned_by=user['id'],
                    workflow_step_id=first_step['id']
                )

    log_audit('form_submission', submission_id, 'SUBMIT', user['id'])

    return jsonify({'success': True, 'message': 'Form submitted successfully'})


@form_bp.route('/api/submissions/<int:submission_id>/approve', methods=['POST'])
@require_login
def api_approve_submission(submission_id):
    """API: Approve a submission."""
    user = get_current_user()
    data = request.get_json() or {}
    comment = data.get('comment', '')

    submission = get_submission(submission_id)
    if not submission:
        return jsonify({'success': False, 'message': 'Submission not found'}), 404

    # Record approval
    workflow_step_id = submission.get('current_step_id') or 0
    add_approval(submission_id, workflow_step_id, user['id'], 'approve', comment)

    # Update status
    update_submission_status(submission_id, SUBMISSION_STATUS_APPROVED, user['id'], reason=comment)

    # Notify submitter
    create_form_notification(
        submission_id=submission_id,
        notification_type='approval',
        title='Form Approved',
        message=f"Your form {submission['submission_number']} has been approved.",
        recipient_id=submission['submitted_by']
    )

    log_audit('form_submission', submission_id, 'APPROVE', user['id'], notes=comment)

    return jsonify({'success': True, 'message': 'Form approved successfully'})


@form_bp.route('/api/submissions/<int:submission_id>/reject', methods=['POST'])
@require_login
def api_reject_submission(submission_id):
    """API: Reject a submission."""
    user = get_current_user()
    data = request.get_json() or {}
    comment = data.get('comment', '')

    if not comment:
        return jsonify({'success': False, 'message': 'Reason is required'}), 400

    submission = get_submission(submission_id)
    if not submission:
        return jsonify({'success': False, 'message': 'Submission not found'}), 404

    # Record rejection
    workflow_step_id = submission.get('current_step_id') or 0
    add_approval(submission_id, workflow_step_id, user['id'], 'reject', comment)

    # Update status
    update_submission_status(submission_id, SUBMISSION_STATUS_REJECTED, user['id'], reason=comment)

    # Notify submitter
    create_form_notification(
        submission_id=submission_id,
        notification_type='rejection',
        title='Form Rejected',
        message=f"Your form {submission['submission_number']} has been rejected. Reason: {comment}",
        recipient_id=submission['submitted_by']
    )

    log_audit('form_submission', submission_id, 'REJECT', user['id'], notes=comment)

    return jsonify({'success': True, 'message': 'Form rejected'})


@form_bp.route('/api/submissions/<int:submission_id>/return', methods=['POST'])
@require_login
def api_return_submission(submission_id):
    """API: Return a submission for correction."""
    user = get_current_user()
    data = request.get_json() or {}
    comment = data.get('comment', '')

    if not comment:
        return jsonify({'success': False, 'message': 'Reason is required'}), 400

    submission = get_submission(submission_id)
    if not submission:
        return jsonify({'success': False, 'message': 'Submission not found'}), 404

    # Update status
    update_submission_status(submission_id, SUBMISSION_STATUS_RETURNED, user['id'], reason=comment)

    # Notify submitter
    create_form_notification(
        submission_id=submission_id,
        notification_type='return',
        title='Form Returned',
        message=f"Your form {submission['submission_number']} has been returned. Reason: {comment}",
        recipient_id=submission['submitted_by']
    )

    log_audit('form_submission', submission_id, 'RETURN', user['id'], notes=comment)

    return jsonify({'success': True, 'message': 'Form returned for correction'})


@form_bp.route('/api/submissions/<int:submission_id>/delegate', methods=['POST'])
@require_login
def api_delegate_submission(submission_id):
    """API: Delegate a submission to another user."""
    user = get_current_user()
    data = request.get_json() or {}
    delegate_to = data.get('delegate_to')
    reason = data.get('reason', '')

    if not delegate_to:
        return jsonify({'success': False, 'message': 'Delegate user is required'}), 400

    submission = get_submission(submission_id)
    if not submission:
        return jsonify({'success': False, 'message': 'Submission not found'}), 404

    # Update assignment
    with get_db_context() as db:
        db.execute("""
            UPDATE form_assignments
            SET is_active = 0, reassigned_to = ?, reassigned_at = CURRENT_TIMESTAMP
            WHERE submission_id = ? AND is_active = 1
        """, (delegate_to, submission_id))
        db.commit()

    # Create new assignment
    assign_submission(
        submission_id=submission_id,
        assigned_to=delegate_to,
        assigned_by=user['id'],
        workflow_step_id=submission.get('current_step_id'),
        notes=f"Delegated: {reason}" if reason else "Delegated"
    )

    log_audit('form_submission', submission_id, 'DELEGATE', user['id'],
              notes=f"Delegated to user {delegate_to}: {reason}")

    return jsonify({'success': True, 'message': 'Submission delegated'})


@form_bp.route('/api/submissions/<int:submission_id>/escalate', methods=['POST'])
@require_login
def api_escalate_submission(submission_id):
    """API: Escalate a submission."""
    user = get_current_user()
    data = request.get_json() or {}
    escalate_to = data.get('escalate_to')
    reason = data.get('reason', '')

    submission = get_submission(submission_id)
    if not submission:
        return jsonify({'success': False, 'message': 'Submission not found'}), 404

    # Find escalation step in workflow
    workflow = get_template_workflow(submission['template_id'])
    escalation_step = None
    if workflow:
        steps = get_workflow_steps(workflow['id'])
        for step in steps:
            if step.get('escalation_step_id'):
                escalation_step = step
                break

    if escalation_step and escalation_step['approver_id']:
        escalate_to = escalation_step['approver_id']

    if escalate_to:
        assign_submission(
            submission_id=submission_id,
            assigned_to=escalate_to,
            assigned_by=user['id'],
            workflow_step_id=escalation_step['id'] if escalation_step else None,
            priority='high',
            notes=f"Escalated: {reason}" if reason else "Escalated"
        )

        # Update status to indicate escalation
        update_submission_status(submission_id, 'escalated', user['id'], reason=f"Escalated: {reason}")

    log_audit('form_submission', submission_id, 'ESCALATE', user['id'],
              notes=f"Escalated: {reason}")

    return jsonify({'success': True, 'message': 'Submission escalated'})


@form_bp.route('/api/submissions/<int:submission_id>/comments')
@require_login
def api_get_comments(submission_id):
    """API: Get all comments for a submission."""
    comments = get_submission_comments(submission_id, include_internal=True)
    return jsonify({'success': True, 'comments': comments})


@form_bp.route('/api/submissions/<int:submission_id>/comments', methods=['POST'])
@require_login
def api_add_comment(submission_id):
    """API: Add a comment to a submission."""
    user = get_current_user()
    data = request.get_json() or {}

    comment_id = add_comment(
        submission_id=submission_id,
        comment_text=data.get('comment_text', ''),
        created_by=user['id'],
        comment_type=data.get('comment_type', 'general'),
        is_internal=data.get('is_internal', 0),
        mentioned_users=data.get('mentioned_users')
    )

    log_audit('form_comment', comment_id, 'CREATE', user['id'])

    return jsonify({'success': True, 'comment_id': comment_id})


@form_bp.route('/api/submissions/<int:submission_id>/history')
@require_login
def api_get_history(submission_id):
    """API: Get status history for a submission."""
    history = get_submission_history(submission_id)
    return jsonify({'success': True, 'history': history})


@form_bp.route('/api/submissions/<int:submission_id>/attachments')
@require_login
def api_get_attachments(submission_id):
    """API: Get all attachments for a submission."""
    attachments = get_submission_attachments(submission_id)
    return jsonify({'success': True, 'attachments': attachments})


@form_bp.route('/api/submissions/<int:submission_id>/attachments', methods=['POST'])
@require_login
def api_add_attachment(submission_id):
    """API: Add an attachment to a submission."""
    user = get_current_user()

    if 'file' not in request.files:
        return jsonify({'success': False, 'message': 'No file provided'}), 400

    file = request.files['file']
    if file.filename == '':
        return jsonify({'success': False, 'message': 'No file selected'}), 400

    # Save file
    filename = secure_filename(file.filename)
    upload_dir = os.path.join('static', 'uploads', 'forms', str(submission_id))
    os.makedirs(upload_dir, exist_ok=True)
    filepath = os.path.join(upload_dir, filename)
    file.save(filepath)

    # Get file size
    file.seek(0, 2)
    file_size = file.tell()
    file.seek(0)

    attachment_id = add_attachment(
        submission_id=submission_id,
        file_name=filename,
        file_path=filepath,
        file_type=file.content_type,
        file_size=file_size,
        uploaded_by=user['id'],
        section_id=request.form.get('section_id'),
        attachment_type=request.form.get('attachment_type', 'general')
    )

    log_audit('form_attachment', attachment_id, 'UPLOAD', user['id'],
              notes=f"Attachment added: {filename}")

    return jsonify({'success': True, 'attachment_id': attachment_id})


# ============================================================================
# AJAX/API HELPERS FOR TEMPLATE BUILDER
# ============================================================================

@form_bp.route('/api/sections/reorder', methods=['POST'])
@require_login
def api_reorder_sections():
    """API: Reorder sections in a template."""
    data = request.get_json() or {}
    section_orders = data.get('orders', [])

    with get_db_context() as db:
        for section_id, new_order in section_orders:
            db.execute(
                "UPDATE form_sections SET section_order = ? WHERE id = ?",
                (new_order, section_id)
            )
        db.commit()

    return jsonify({'success': True})


@form_bp.route('/api/fields/reorder', methods=['POST'])
@require_login
def api_reorder_fields():
    """API: Reorder fields in a section."""
    data = request.get_json() or {}
    field_orders = data.get('orders', [])

    with get_db_context() as db:
        for field_id, new_order in field_orders:
            db.execute(
                "UPDATE form_fields SET field_order = ? WHERE id = ?",
                (new_order, field_id)
            )
        db.commit()

    return jsonify({'success': True})


# ============================================================================
# REGISTRATION FUNCTION
# ============================================================================

# =============================================================================
# API EXPORT ENDPOINTS
# =============================================================================

@form_bp.route('/api/export/<export_type>', methods=['GET', 'POST'])
@form_bp.route('/api/export/<data_type>/<export_type>', methods=['GET', 'POST'])
@require_login
def api_form_export(export_type, data_type=None):
    """Export form builder data in all 20 formats."""
    if export_type not in FORM_EXPORT_TYPES:
        return jsonify({
            'error': f'Invalid export type. Valid types: {FORM_EXPORT_TYPES}'
        }), 400

    company_id = session.get('company_id', 0)

    # Determine data type from URL or default
    if data_type is None:
        data_type = request.args.get('type', 'templates')

    db = get_db_context()
    try:
        # Get data based on type
        if data_type == 'templates':
            data = db.execute("""
                SELECT * FROM form_templates
                WHERE company_id = ?
                ORDER BY created_at DESC
                LIMIT 5000
            """, (company_id,)).fetchall()
            data = [dict(row) for row in data]
            columns = FORM_EXPORT_COLUMNS['templates']
            title = 'Form Templates'
        elif data_type == 'submissions':
            data = db.execute("""
                SELECT s.*, t.template_name
                FROM form_submissions s
                LEFT JOIN form_templates t ON s.template_id = t.id
                WHERE s.company_id = ?
                ORDER BY s.submitted_at DESC
                LIMIT 5000
            """, (company_id,)).fetchall()
            data = [dict(row) for row in data]
            columns = FORM_EXPORT_COLUMNS['submissions']
            title = 'Form Submissions'
        elif data_type == 'drafts':
            data = db.execute("""
                SELECT d.*, t.template_name
                FROM form_drafts d
                LEFT JOIN form_templates t ON d.template_id = t.id
                WHERE d.company_id = ?
                ORDER BY d.last_modified DESC
                LIMIT 5000
            """, (company_id,)).fetchall()
            data = [dict(row) for row in data]
            columns = FORM_EXPORT_COLUMNS['drafts']
            title = 'Form Drafts'
        elif data_type == 'approvals':
            data = db.execute("""
                SELECT a.*, s.submission_id as ref_id
                FROM form_approvals a
                LEFT JOIN form_submissions s ON a.submission_id = s.id
                WHERE a.company_id = ?
                ORDER BY a.action_at DESC
                LIMIT 5000
            """, (company_id,)).fetchall()
            data = [dict(row) for row in data]
            columns = FORM_EXPORT_COLUMNS['approvals']
            title = 'Form Approvals'
        else:
            return jsonify({'error': f'Data type {data_type} not supported'}), 400

        filename = f'form_{data_type}_{datetime.now().strftime("%Y%m%d")}'

        return send_export_response(data, export_type, filename, columns, title)
    finally:
        db.close()


@form_bp.route('/api/export/list')
@require_login
def list_form_export_types():
    """List available export types for form builder module."""
    return jsonify({
        'module': 'forms',
        'data_types': list(FORM_EXPORT_COLUMNS.keys()),
        'export_types': [{'type': t} for t in FORM_EXPORT_TYPES]
    })


def register_form_routes(app):
    """Register form routes with the Flask app."""
    initialize_form_builder_tables()
    app.register_blueprint(form_bp)
